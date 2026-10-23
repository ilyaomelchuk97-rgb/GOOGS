#!/usr/bin/env python3
"""Локальное веб-приложение учёта работ. Python stdlib + SQLite; openpyxl нужен только для импорта/экспорта XLSX."""
import base64
import binascii
import calendar
import csv
import ipaddress
import datetime as dt
import hashlib
import hmac
import io
import json
import mimetypes
import math
import os
import re
import secrets
import sqlite3
import threading
import traceback
import time
import urllib.request
import urllib.error
import backup_tools
import d1_store
import vacation
import sotipb
from http import HTTPStatus
from http.cookies import SimpleCookie
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from urllib.parse import parse_qs, urlparse
from zoneinfo import ZoneInfo

ROOT = Path(__file__).resolve().parent
DB = ROOT / 'worktrack.sqlite3'
BACKUP_DIR = ROOT / 'backups'
DB_LOCK = threading.RLock()
CALL_LAST_CLEANUP = 0.0
BUILD_ID = '20261023-51'  # Public /health marker to verify which build Render actually serves.
# New effects reuse the existing CHECK(kind IN ('people','speech')) table safely.
# This keeps old production D1/SQLite backups and schema compatible.
PRANK_EFFECT_PREFIX = '\x1eFORMA_EFFECT_V1:'
PRANK_NEW_KINDS = ('cat','achievement','parade','plane','stamp','duck','stickers','bonus')
# Background tabs may be throttled to ~1-minute timers; 15 seconds was too short.
PRESENCE_WINDOW = dt.timedelta(seconds=120)
TZ = ZoneInfo('Europe/Minsk')
TODAY = lambda: dt.datetime.now(TZ).date()
MONTH_PATTERN = re.compile(r'^\d{4}-(0[1-9]|1[0-2])$')
USERNAME_PATTERN = re.compile(r'^[a-zA-Z0-9_.-]{3,40}$')

def connect():
    backend=os.environ.get('DATABASE_BACKEND','sqlite').lower()
    if os.environ.get('RENDER')=='true' and backend!='d1':
        raise RuntimeError('Render requires DATABASE_BACKEND=d1; refusing ephemeral SQLite')
    if backend=='d1': return d1_store.D1Connection()
    if backend!='sqlite': raise RuntimeError('Unsupported DATABASE_BACKEND')
    con = sqlite3.connect(DB, timeout=15)
    con.row_factory = sqlite3.Row
    con.execute('PRAGMA foreign_keys=ON')
    con.execute('PRAGMA busy_timeout=15000')
    return con

def init_db():
    with connect() as db:
        db.executescript('''
        PRAGMA journal_mode=WAL;
        CREATE TABLE IF NOT EXISTS users (
          id INTEGER PRIMARY KEY, name TEXT NOT NULL, username TEXT NOT NULL UNIQUE COLLATE NOCASE,
          password_hash TEXT NOT NULL, role TEXT NOT NULL CHECK(role IN ('admin','employee')),
          is_staff INTEGER NOT NULL DEFAULT 1 CHECK(is_staff IN (0,1)),
          active INTEGER NOT NULL DEFAULT 1, created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP);
        CREATE TABLE IF NOT EXISTS tasks (
          id INTEGER PRIMARY KEY, title TEXT NOT NULL UNIQUE COLLATE NOCASE, unit TEXT NOT NULL DEFAULT 'шт.',
          norm REAL NOT NULL DEFAULT 0, category TEXT NOT NULL DEFAULT 'Другое', active INTEGER NOT NULL DEFAULT 1,
          created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP);
        CREATE TABLE IF NOT EXISTS entries (
          id INTEGER PRIMARY KEY, user_id INTEGER NOT NULL REFERENCES users(id), task_id INTEGER NOT NULL REFERENCES tasks(id),
          work_date TEXT NOT NULL, quantity REAL NOT NULL CHECK(quantity > 0), note TEXT NOT NULL DEFAULT '',
          created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP, updated_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,
          UNIQUE(user_id,task_id,work_date));
        CREATE INDEX IF NOT EXISTS idx_entries_date ON entries(work_date);
        CREATE TABLE IF NOT EXISTS cell_comments (
          user_id INTEGER NOT NULL REFERENCES users(id), task_id INTEGER NOT NULL REFERENCES tasks(id),
          work_date TEXT NOT NULL, body TEXT NOT NULL,
          updated_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,
          PRIMARY KEY(user_id,task_id,work_date));
        CREATE INDEX IF NOT EXISTS idx_cell_comments_date ON cell_comments(work_date);
        CREATE TABLE IF NOT EXISTS attendance (
          id INTEGER PRIMARY KEY, user_id INTEGER NOT NULL REFERENCES users(id), day TEXT NOT NULL,
          status TEXT NOT NULL CHECK(status IN ('present','absent')), reason TEXT NOT NULL DEFAULT '',
          updated_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP, UNIQUE(user_id,day));
        CREATE TABLE IF NOT EXISTS daily_hours (
          day TEXT PRIMARY KEY, hours REAL NOT NULL CHECK(hours BETWEEN 0 AND 24),
          updated_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP);
        CREATE TABLE IF NOT EXISTS personal_hours (
          user_id INTEGER NOT NULL REFERENCES users(id), day TEXT NOT NULL,
          hours REAL NOT NULL CHECK(hours BETWEEN 0 AND 24),
          updated_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,
          PRIMARY KEY(user_id,day));
        CREATE TABLE IF NOT EXISTS group_chats (
         id INTEGER PRIMARY KEY, title TEXT NOT NULL, created_by INTEGER NOT NULL REFERENCES users(id),
         created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP);
        CREATE TABLE IF NOT EXISTS group_members (
         group_id INTEGER NOT NULL REFERENCES group_chats(id), user_id INTEGER NOT NULL REFERENCES users(id),
         joined_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP, PRIMARY KEY(group_id,user_id));
        CREATE TABLE IF NOT EXISTS group_messages (
         id INTEGER PRIMARY KEY, group_id INTEGER NOT NULL REFERENCES group_chats(id),
         sender_id INTEGER NOT NULL REFERENCES users(id), body TEXT NOT NULL,
         style_json TEXT NOT NULL DEFAULT '{}', created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP);
        CREATE INDEX IF NOT EXISTS idx_group_messages_room ON group_messages(group_id,id);
        CREATE TABLE IF NOT EXISTS call_sessions (
         id INTEGER PRIMARY KEY, room_key TEXT NOT NULL, mode TEXT NOT NULL,
         created_by INTEGER NOT NULL, state TEXT NOT NULL DEFAULT 'active',
         created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP, updated_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP);
        CREATE TABLE IF NOT EXISTS call_participants (
         call_id INTEGER NOT NULL REFERENCES call_sessions(id), user_id INTEGER NOT NULL,
         connection_id TEXT NOT NULL, last_seen TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,
         PRIMARY KEY(call_id,connection_id));
        CREATE TABLE IF NOT EXISTS call_signals (
         id INTEGER PRIMARY KEY, call_id INTEGER NOT NULL REFERENCES call_sessions(id),
         sender_conn TEXT NOT NULL, target_conn TEXT NOT NULL, signal_json TEXT NOT NULL,
         created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP);
        CREATE TABLE IF NOT EXISTS chat_messages (
          id INTEGER PRIMARY KEY, sender_id INTEGER NOT NULL REFERENCES users(id),
          recipient_id INTEGER REFERENCES users(id), body TEXT NOT NULL,
          created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,
          style_json TEXT NOT NULL DEFAULT '{}');
        CREATE INDEX IF NOT EXISTS idx_chat_recipient ON chat_messages(recipient_id,id);
        CREATE INDEX IF NOT EXISTS idx_chat_sender ON chat_messages(sender_id,id);
        CREATE TABLE IF NOT EXISTS chat_reads (
          user_id INTEGER NOT NULL REFERENCES users(id), room_key TEXT NOT NULL,
          last_read_id INTEGER NOT NULL DEFAULT 0, PRIMARY KEY(user_id,room_key));
        CREATE TABLE IF NOT EXISTS chat_themes (
          room_key TEXT PRIMARY KEY, theme_json TEXT NOT NULL,
          updated_by INTEGER NOT NULL REFERENCES users(id),
          updated_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP);
CREATE TABLE IF NOT EXISTS vacation_schedule (
 user_id INTEGER NOT NULL REFERENCES users(id),
 year INTEGER NOT NULL CHECK(year BETWEEN 2020 AND 2100),
 month INTEGER NOT NULL CHECK(month BETWEEN 1 AND 12),
 start_date TEXT NOT NULL, end_date TEXT NOT NULL,
 material_help INTEGER NOT NULL DEFAULT 0 CHECK(material_help IN (0,1)),
 updated_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,
 PRIMARY KEY(user_id,year,month));
        CREATE TABLE IF NOT EXISTS chat_pins (
         user_id INTEGER NOT NULL REFERENCES users(id), room_key TEXT NOT NULL,
         created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,
         PRIMARY KEY(user_id,room_key));
        CREATE TABLE IF NOT EXISTS sotipb_runs (
          id INTEGER PRIMARY KEY, uploader_id INTEGER NOT NULL REFERENCES users(id),
          source_name TEXT NOT NULL, source_size INTEGER NOT NULL, result_size INTEGER NOT NULL,
          row_count INTEGER NOT NULL, group_count INTEGER NOT NULL,
          warnings_json TEXT NOT NULL DEFAULT '[]', source_sha256 TEXT NOT NULL,
          result_sha256 TEXT NOT NULL, state TEXT NOT NULL DEFAULT 'pending',
          created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP);
        CREATE TABLE IF NOT EXISTS sotipb_parts (
          run_id INTEGER NOT NULL REFERENCES sotipb_runs(id),
          kind TEXT NOT NULL CHECK(kind IN ('source','result')),
          part_no INTEGER NOT NULL, content_b64 TEXT NOT NULL,
          PRIMARY KEY(run_id,kind,part_no));
        CREATE TABLE IF NOT EXISTS backup_monitor (
         source TEXT PRIMARY KEY, last_check TEXT NOT NULL DEFAULT '',
         last_saved TEXT NOT NULL DEFAULT '', file_name TEXT NOT NULL DEFAULT '',
         file_size INTEGER NOT NULL DEFAULT 0, signature TEXT NOT NULL DEFAULT '',
         state TEXT NOT NULL DEFAULT 'unknown', detail TEXT NOT NULL DEFAULT '');
        CREATE TABLE IF NOT EXISTS restore_stage (
          transfer_id TEXT NOT NULL, table_name TEXT NOT NULL, chunk_no INTEGER NOT NULL,
          payload TEXT NOT NULL, created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,
          PRIMARY KEY(transfer_id,table_name,chunk_no));
        CREATE TABLE IF NOT EXISTS user_avatars (
          user_id INTEGER PRIMARY KEY REFERENCES users(id), image_b64 TEXT NOT NULL DEFAULT '',
          mime TEXT NOT NULL DEFAULT '', image_url TEXT NOT NULL DEFAULT '', revision INTEGER NOT NULL DEFAULT 1);
        CREATE TABLE IF NOT EXISTS sessions (
          token_hash TEXT PRIMARY KEY, user_id INTEGER NOT NULL REFERENCES users(id),
          expires_at TEXT NOT NULL);
        CREATE TABLE IF NOT EXISTS activity (
          id INTEGER PRIMARY KEY, actor_id INTEGER REFERENCES users(id), action TEXT NOT NULL,
          detail TEXT NOT NULL, created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP);
        CREATE TABLE IF NOT EXISTS prank_presence (
          session_hash TEXT PRIMARY KEY, user_id INTEGER NOT NULL REFERENCES users(id),
          last_seen TEXT NOT NULL);
        CREATE INDEX IF NOT EXISTS idx_prank_presence_user ON prank_presence(user_id,last_seen);
        CREATE TABLE IF NOT EXISTS prank_events (
          id INTEGER PRIMARY KEY, target_id INTEGER NOT NULL REFERENCES users(id),
          kind TEXT NOT NULL CHECK(kind IN ('people','speech')),
          body TEXT NOT NULL DEFAULT '', created_at TEXT NOT NULL);
        CREATE INDEX IF NOT EXISTS idx_prank_events_target ON prank_events(target_id,id);
        ''')
        if 'style_json' not in [col[1] for col in db.execute('PRAGMA table_info(chat_messages)')]:
            db.execute("ALTER TABLE chat_messages ADD COLUMN style_json TEXT NOT NULL DEFAULT '{}'")
        # Право администратора отдельно от принадлежности к сотрудникам: повышение
        # не удаляет личные работы и табель. Сводный отчёт отдельно фильтрует роль.
        if 'is_staff' not in [col[1] for col in db.execute('PRAGMA table_info(users)')]:
            db.execute('ALTER TABLE users ADD COLUMN is_staff INTEGER NOT NULL DEFAULT 1')
        db.execute("UPDATE users SET is_staff=0 WHERE username='admin' AND role='admin' AND is_staff<>0")
        # Исправляем старые повышения, при которых работнику вместе с ролью admin
        # ошибочно ставили is_staff=0: это скрывало его из табеля и общих графиков.
        # Сохраняем одну исходную нештатную учётную запись администратора.
        primary=db.execute('''SELECT id FROM users WHERE role='admin' AND is_staff=0
            ORDER BY CASE WHEN username='admin' THEN 0 ELSE 1 END,created_at,id LIMIT 1''').fetchone()
        if primary:
            db.execute('UPDATE users SET is_staff=1 WHERE role=? AND is_staff=0 AND id<>?',
                       ('admin',primary['id']))
        # Перенос комментариев из старых записей при обновлении существующей базы.
        db.execute('''INSERT OR IGNORE INTO cell_comments(user_id,task_id,work_date,body)
            SELECT user_id,task_id,work_date,note FROM entries WHERE note<>?''',('',))

def hashed_password(password, salt=None):
    salt = salt or secrets.token_bytes(16)
    digest = hashlib.pbkdf2_hmac('sha256', password.encode(), salt, 260000)
    return salt.hex() + ':' + digest.hex()

def valid_password(stored, password):
    try:
        salt, digest = stored.split(':')
        actual = hashed_password(password, bytes.fromhex(salt)).split(':')[1]
        return hmac.compare_digest(actual, digest)
    except (ValueError, TypeError): return False

def json_rows(rows): return [dict(r) for r in rows]
CHAT_STYLE_KEYS = {'bubble', 'bubble_alpha', 'glow', 'glow_color', 'glow_alpha',
                   'border_effect', 'flower', 'branch', 'logo'}
def valid_chat_style(style):
    """Strict finite color, opacity and ornament indexes; never store arbitrary CSS."""
    if style is None: return {}
    if not isinstance(style,dict) or set(style)-CHAT_STYLE_KEYS:
        raise APIError(400,'Некорректное оформление чата')
    clean={}
    for key in ('bubble','glow_color'):
        if key in style:
            color=style[key]
            if not isinstance(color,str) or not re.fullmatch(r'#[0-9a-fA-F]{6}',color):
                raise APIError(400,'Цвет должен быть в формате #RRGGBB')
            clean[key]=color.lower()
    if 'glow' in style:
        if type(style['glow']) is not bool:raise APIError(400,'Некорректная настройка подсветки')
        clean['glow']=style['glow']
    for key in ('bubble_alpha','glow_alpha'):
        if key in style:
            value=style[key]
            if type(value) is not int or not 0<=value<=100:
                raise APIError(400,'Прозрачность: от 0 до 100 процентов')
            clean[key]=value
    for key in ('border_effect','flower','branch','logo'):
        if key in style:
            value=style[key]
            limit=26 if key=='border_effect' else 20
            if type(value) is not int or not 0<=value<=limit:
                raise APIError(400,f'Номер украшения: от 0 до {limit}')
            clean[key]=value
    return clean

def checked_group(db,user,room):
    if not isinstance(room,str) or not re.fullmatch(r'group:[1-9]\d*',room):
        raise APIError(400,'Некорректная групповая беседа')
    group_id=int(room.split(':')[1])
    row=db.execute('''SELECT g.id,g.title,g.created_by FROM group_chats g
        JOIN group_members gm ON gm.group_id=g.id
        WHERE g.id=? AND gm.user_id=?''',(group_id,user['id'])).fetchone()
    if not row:raise APIError(403,'Вы не участник этой беседы')
    return row

def call_room_key(db,user,room):
    if room=='general':return 'general'
    if isinstance(room,str) and room.startswith('group:'):
        return 'group:'+str(checked_group(db,user,room)['id'])
    try:peer=int(room)
    except (TypeError,ValueError):raise APIError(400,'Некорректный собеседник')
    if str(peer)!=room or peer<=0 or peer==user['id'] or not db.execute('SELECT 1 FROM users WHERE id=? AND active=1',(peer,)).fetchone():
        raise APIError(400,'Собеседник не найден')
    return 'dm:'+str(min(peer,user['id']))+':'+str(max(peer,user['id']))

def chat_theme_key(db,user,room):
    if room=='general':return 'general'
    if isinstance(room,str) and room.startswith('group:'):return 'group:'+str(checked_group(db,user,room)['id'])
    try: peer=int(room)
    except (TypeError,ValueError):raise APIError(400,'Некорректный собеседник')
    if peer<=0 or str(peer)!=room or peer==user['id'] or not db.execute('SELECT 1 FROM users WHERE id=?',(peer,)).fetchone():
        raise APIError(400,'Собеседник не найден')
    return 'pair:'+str(min(user['id'],peer))+':'+str(max(user['id'],peer))

def livekit_ready():
    return all(os.environ.get(key) for key in ('LIVEKIT_URL','LIVEKIT_API_KEY','LIVEKIT_API_SECRET')) and os.environ['LIVEKIT_URL'].startswith('wss://')

def livekit_jwt(video,subject,name='',lifetime=600):
    now=int(dt.datetime.now(dt.timezone.utc).timestamp())
    b64=lambda raw:base64.urlsafe_b64encode(raw).rstrip(b'=').decode()
    header=b64(b'{"alg":"HS256","typ":"JWT"}')
    payload=b64(json.dumps({'iss':os.environ['LIVEKIT_API_KEY'],'sub':subject,'name':name,
        'iat':now,'nbf':now-5,'exp':now+lifetime,'video':video},separators=(',',':')).encode())
    signature=b64(hmac.new(os.environ['LIVEKIT_API_SECRET'].encode(),(header+'.'+payload).encode(),hashlib.sha256).digest())
    return header+'.'+payload+'.'+signature

def livekit_token(call,user,connection):
    return livekit_jwt({'roomJoin':True,'room':'forma-call-'+str(call['id']),
        'canPublish':True,'canSubscribe':True,'canPublishData':False},connection,user['name'])

def livekit_evict(call_id,connection):
    """Revoke LiveKit Cloud token and disconnect active participant on membership removal."""
    if not livekit_ready():raise APIError(503,'Нельзя изменить состав во время LiveKit-звонка: ключи LiveKit недоступны')
    room='forma-call-'+str(call_id)
    token=livekit_jwt({'roomAdmin':True,'room':room},'forma-server',lifetime=60)
    data=json.dumps({'room':room,'identity':connection}).encode()
    url='https://'+os.environ['LIVEKIT_URL'][6:].rstrip('/')+'/twirp/livekit.RoomService/RemoveParticipant'
    request=urllib.request.Request(url,data,{'Authorization':'Bearer '+token,'Content-Type':'application/json'},method='POST')
    try:
        with urllib.request.urlopen(request,timeout=5) as response:response.read(1024)
    except urllib.error.HTTPError as exc:
        # If the participant has left, LiveKit Cloud still revokes the token.
        if exc.code!=404:raise APIError(503,'LiveKit не подтвердил исключение участника; состав не изменён') from exc
    except (urllib.error.URLError,TimeoutError,OSError) as exc:
        raise APIError(503,'LiveKit недоступен: участник остаётся в беседе, повторите позже') from exc

def log(db, user_id, action, detail):
    db.execute('INSERT INTO activity(actor_id,action,detail) VALUES(?,?,?)',(user_id,action,detail[:350]))

def date_ok(value, future=False):
    try:
        day = dt.date.fromisoformat(value)
        if day.isoformat() != value or day.year < 2020 or day > TODAY()+dt.timedelta(days=366 if future else 0): raise ValueError()
        return day
    except (TypeError, ValueError): raise APIError(400,'Некорректная дата')

def scheduled_hours(day):
    weekday=dt.date.fromisoformat(day).weekday()
    return 8.25 if weekday<4 else 7.0 if weekday==4 else 0.0

def calculate_productivity(month,rows,users,attendance,daily_hours,personal_hours):
    """Как в исходном Excel: ROUND(нормо-часы, 0) / часы по табелю × 100."""
    count=calendar.monthrange(*map(int,month.split('-')))[1]
    by_user={}
    for row in rows:
        by_user[row['user_id']]=by_user.get(row['user_id'],0)+float(row['hours'])
    status_map={(a['user_id'],a['day']):a['status'] for a in attendance}
    global_map={a['day']:float(a['hours']) for a in daily_hours}
    personal_map={(a['user_id'],a['day']):float(a['hours']) for a in personal_hours}
    people=[]
    for user in users:
        uid=user['id'];worked=0.0
        for i in range(1,count+1):
            day=f'{month}-{i:02d}'
            status=status_map.get((uid,day),'present' if dt.date.fromisoformat(day).weekday()<5 else 'off')
            if status=='present':worked+=personal_map.get((uid,day),global_map.get(day,scheduled_hours(day)))
        norm=by_user.get(uid,0.0)
        rounded_norm=math.floor(norm+0.5) # Excel ROUND(...,0), для положительных нормо-часов
        people.append({'user_id':uid,'name':user['name'],'work_hours':round(worked,2),
                       'norm_hours':round(norm,3),'rounded_norm_hours':rounded_norm,
                       'percent':round(rounded_norm/worked*100,2) if worked else None})
    total_work=round(sum(row['work_hours'] for row in people),2)
    total_norm=round(sum(row['norm_hours'] for row in people),3)
    total_rounded=sum(row['rounded_norm_hours'] for row in people)
    overall={'work_hours':total_work,'norm_hours':total_norm,'rounded_norm_hours':total_rounded,
             'percent':round(total_rounded/total_work*100,2) if total_work else None}
    return people,overall

def month_range(month):
    if not MONTH_PATTERN.fullmatch(month or ''): raise APIError(400,'Некорректный месяц')
    year, mon = map(int,month.split('-'))
    if not 2020 <= year <= TODAY().year+1: raise APIError(400,'Месяц вне допустимого диапазона')
    return month+'-01', f'{year}-{mon:02d}-{calendar.monthrange(year,mon)[1]:02d}'

class APIError(Exception):
    def __init__(self, status, message): self.status=status; self.message=message

def check_expected(db, data, table, criteria, params, fields):
    """Compare a client's last seen row with the live row before any mutation.

    Existing old clients without expected remain compatible; new UI sends it on edits.
    The process lock serializes requests on the single Render instance.
    """
    if 'expected' not in data:return
    prior=data['expected']
    if prior is not None and (not isinstance(prior,dict) or set(prior)!=set(fields)):
        raise APIError(400,'Некорректная версия записи')
    row=db.execute('SELECT '+','.join('"'+f+'"' for f in fields)+' FROM '+table+' WHERE '+criteria,params).fetchone()
    current={field:row[field] for field in fields} if row else None
    if current!=prior:
        raise APIError(409,'Данные уже изменены в другой вкладке или другим пользователем. Обновите страницу и проверьте новое значение перед повтором')

class Handler(BaseHTTPRequestHandler):
    server_version = 'Forma/1.0'
    def log_message(self, fmt, *args):
        # Не выводим резервный токен сеанса в журнал HTTP-запросов.
        message=re.sub(r'([?&]forma_ticket=)[^&\s]+',r'\1[redacted]',fmt%args)
        print('%s %s'%(self.address_string(),message),flush=True)
    def send_json(self, data, status=200, cookie=None):
        raw=json.dumps(data,ensure_ascii=False,default=str).encode()
        self.send_response(status); self.send_header('Content-Type','application/json; charset=utf-8')
        self.send_header('Content-Length',str(len(raw))); self.send_header('Cache-Control','no-store')
        self.send_header('X-Content-Type-Options','nosniff'); self.send_header('Referrer-Policy','no-referrer')
        if cookie: self.send_header('Set-Cookie',cookie)
        self.end_headers(); self.wfile.write(raw)
    def send_bytes(self, raw, filename, mime):
        self.send_response(200); self.send_header('Content-Type',mime)
        self.send_header('Content-Disposition',f'attachment; filename="{filename}"')
        self.send_header('Content-Length',str(len(raw))); self.send_header('Cache-Control','no-store'); self.end_headers(); self.wfile.write(raw)
    def auth_token(self):
        # Прокси предпросмотра может блокировать cookie и зарезервировать Authorization.
        # Свой заголовок позволяет работать в iframe без сторонних cookie.
        token=self.headers.get('X-Forma-Session','').strip()
        if token: return token if len(token)<=256 else None
        # Резервный путь: некоторые прокси вырезают даже нестандартные заголовки.
        ticket=parse_qs(urlparse(self.path).query).get('forma_ticket',[''])[0]
        if ticket: return ticket if len(ticket)<=256 else None
        authorization=self.headers.get('Authorization','')
        if authorization.startswith('Bearer '):
            token=authorization[7:].strip()
            return token if len(token)<=256 else None
        cookie=SimpleCookie()
        try: cookie.load(self.headers.get('Cookie',''))
        except Exception: return None
        return cookie['session'].value if 'session' in cookie else None
    def user(self, db):
        token=self.auth_token()
        if not token: return None
        return db.execute('''SELECT u.id,u.name,u.username,u.role,u.is_staff,u.active,
            COALESCE(a.revision,0) avatar_revision,COALESCE(a.image_url,'') avatar_url,
            CASE WHEN a.image_b64<>'' THEN 1 ELSE 0 END avatar_stored
            FROM sessions s JOIN users u ON u.id=s.user_id LEFT JOIN user_avatars a ON a.user_id=u.id
            WHERE s.token_hash=? AND s.expires_at>? AND u.active=1''',
            (hashlib.sha256(token.encode()).hexdigest(),dt.datetime.now(dt.timezone.utc).isoformat())).fetchone()
    def require(self, db, admin=False):
        user=self.user(db)
        if not user:
            print('auth rejected: x-header=%s bearer=%s cookie=%s' %
                  (bool(self.headers.get('X-Forma-Session')),
                   self.headers.get('Authorization','').startswith('Bearer '),
                   bool(self.headers.get('Cookie'))),flush=True)
            raise APIError(401,'Войдите в аккаунт')
        if admin and user['role']!='admin': raise APIError(403,'Требуются права администратора')
        return user
    def body(self,max_bytes=32768):
        if 'application/json' not in self.headers.get('Content-Type',''): raise APIError(415,'Требуется JSON')
        try:
            length=int(self.headers.get('Content-Length','0'))
            if length<1 or length>max_bytes: raise APIError(413,'Неверный размер запроса')
            data=json.loads(self.rfile.read(length))
            if not isinstance(data,dict): raise ValueError()
            return data
        except (ValueError,UnicodeDecodeError): raise APIError(400,'Некорректный JSON')
    def backup_body(self):
        if self.headers.get('Content-Type','').split(';')[0] not in ('application/octet-stream','application/zip'):
            raise APIError(415,'Загрузите ZIP-файл резервной копии')
        try: size=int(self.headers.get('Content-Length','0'))
        except ValueError: raise APIError(400,'Некорректный размер файла')
        if not 0<size<=backup_tools.MAX_ARCHIVE: raise APIError(413,'Архив слишком большой (максимум 25 МБ)')
        return self.rfile.read(size)
    def setup_gate(self, db):
        # This is a one-time, opt-in setup path. It must not be available to
        # ordinary employees or against an already-populated D1 database.
        if not isinstance(db,d1_store.D1Connection): raise APIError(404,'Страница недоступна')
        key=os.environ.get('INITIAL_IMPORT_TOKEN','')
        if len(key)<32: raise APIError(404,'Первоначальная загрузка не включена')
        if any(db.execute('SELECT COUNT(*) FROM "'+table+'"').fetchone()[0]
               for table in d1_store.TABLES):
            raise APIError(409,'База уже содержит данные. Повторная загрузка запрещена')
        return key
    def setup_page(self):
        with connect() as db: self.setup_gate(db)
        data=(ROOT/'web'/'setup.html').read_bytes()
        self.send_response(200)
        self.send_header('Content-Type','text/html; charset=utf-8')
        self.send_header('Content-Length',str(len(data)))
        self.send_header('Cache-Control','no-store')
        self.send_header('X-Content-Type-Options','nosniff')
        self.end_headers(); self.wfile.write(data)
    def setup_import(self):
        with connect() as db:
            expected=self.setup_gate(db)
            supplied=self.headers.get('X-Forma-Setup-Key','')
            if len(supplied)>256 or not hmac.compare_digest(supplied,expected):
                raise APIError(403,'Неверный одноразовый ключ загрузки')
            password=self.headers.get('X-Forma-New-Password','')
            if not re.fullmatch(r'[\x21-\x7e]{14,128}',password):
                raise APIError(400,'Новый пароль администратора: 14–128 латинских символов, цифр и знаков, без пробелов')
            raw=self.backup_body()
            try: data,manifest,details=backup_tools.inspect(raw)
            except backup_tools.BackupError as exc: raise APIError(400,str(exc))
            with sqlite3.connect(':memory:') as source:
                source.deserialize(data)
                primary=source.execute('SELECT id,username FROM users WHERE role="admin" AND active=1 AND is_staff=0 LIMIT 1').fetchone()
                if not primary: raise APIError(400,'В архиве нет основного администратора')
                source.execute('UPDATE users SET password_hash=? WHERE id=?',(hashed_password(password),primary[0]))
                source.commit()
                d1_store.replace_atomic(db,source)
                expected_fingerprint=backup_tools.fingerprint(source)
            # The uploaded ZIP is intentionally never written to Render's disk.
            if (backup_tools.summary(db)['counts']!=details['counts'] or
                backup_tools.fingerprint(db)!=expected_fingerprint):
                raise APIError(503,'Данные отправлены в D1, но проверка не прошла. Не загружайте ZIP снова; обратитесь за помощью')
            return self.send_json({'ok':True,'username':primary[1],'counts':details['counts'],
                                   'first_work_date':details['first_work_date'],
                                   'last_work_date':details['last_work_date']})
    def serve_static(self, path):
        if path=='/': path='/index.html'
        if path not in ('/index.html','/styles.css','/main.js','/calls.js','/livekit-client.umd.min.js','/courier-cat.png','/courier-cat-walk.png',
                            '/chat-flower.png','/chat-branch.png','/mingas-official-logo.webp',
                            '/effect-snake.png','/effect-comet.png','/effect-firefly.png',
                            '/effect-electric.png','/effect-rain.png','/effect-aurora.png'): raise APIError(404,'Страница не найдена')
        file=ROOT/'web'/path[1:]
        data=file.read_bytes(); mime=mimetypes.guess_type(file.name)[0] or 'application/octet-stream'
        content_type=mime if mime.startswith('image/') else mime+'; charset=utf-8'
        self.send_response(200); self.send_header('Content-Type',content_type); self.send_header('Content-Length',str(len(data)))
        self.send_header('X-Content-Type-Options','nosniff'); self.send_header('Referrer-Policy','no-referrer')
        self.send_header('Cache-Control','no-store, max-age=0, must-revalidate')
        self.send_header('Pragma','no-cache')
        self.end_headers(); self.wfile.write(data)
    def do_GET(self): self.dispatch('GET')
    def do_POST(self): self.dispatch('POST')
    def do_PATCH(self): self.dispatch('PATCH')
    def do_DELETE(self): self.dispatch('DELETE')
    def dispatch(self, method):
        # Восстановление БД не должно пересечься с чтением или записью другого запроса.
        with DB_LOCK: self._dispatch(method)
    def _dispatch(self, method):
        parsed=urlparse(self.path); path=parsed.path; q=parse_qs(parsed.query)
        try:
            if path=='/health' and method=='GET':
                return self.send_json({'ok':True,'service':'Forma','backend':os.environ.get('DATABASE_BACKEND','sqlite'),'build':BUILD_ID})
            if path=='/setup' and method=='GET':
                return self.setup_page()
            if path=='/setup/import' and method=='POST':
                return self.setup_import()
            if path=='/internal/backup/report' and method=='POST':
                expected=os.environ.get('AUTOMATED_BACKUP_TOKEN','')
                supplied=self.headers.get('X-Forma-Backup-Key','')
                if len(expected)<32 or len(supplied)>256 or not hmac.compare_digest(supplied,expected):
                    raise APIError(403,'Доступ запрещён')
                data=self.body(4096)
                state=data.get('state');name=data.get('file_name','');size=data.get('file_size',0)
                signature=data.get('signature','');detail=str(data.get('detail',''))[:180]
                if state not in ('saved','unchanged','error') or not isinstance(name,str) or len(name)>180 or not isinstance(size,int) or size<0 or size>100000000 or not isinstance(signature,str) or (signature and not re.fullmatch('[a-f0-9]{64}',signature)):
                    raise APIError(400,'Некорректный отчёт о резервной копии')
                with connect() as db:
                    db.execute('''INSERT INTO backup_monitor(source,last_check,last_saved,file_name,file_size,signature,state,detail)
                        VALUES('google',CURRENT_TIMESTAMP,CASE WHEN ?='saved' THEN CURRENT_TIMESTAMP ELSE '' END,?,?,?,?,?)
                        ON CONFLICT(source) DO UPDATE SET last_check=CURRENT_TIMESTAMP,
                        last_saved=CASE WHEN excluded.state='saved' THEN CURRENT_TIMESTAMP ELSE backup_monitor.last_saved END,
                        file_name=CASE WHEN excluded.state='saved' THEN excluded.file_name ELSE backup_monitor.file_name END,
                        file_size=CASE WHEN excluded.state='saved' THEN excluded.file_size ELSE backup_monitor.file_size END,
                        signature=CASE WHEN excluded.state='saved' THEN excluded.signature ELSE backup_monitor.signature END,
                        state=excluded.state,detail=excluded.detail''',
                        (state,name,size,signature,state,detail))
                    db.commit()
                return self.send_json({'ok':True})
            if path=='/internal/backup' and method=='POST':
                # Dedicated read-only credential for Google Drive backups: never give
                # Apps Script an administrator password or an interactive session.
                expected=os.environ.get('AUTOMATED_BACKUP_TOKEN','')
                supplied=self.headers.get('X-Forma-Backup-Key','')
                if len(expected)<32 or len(supplied)>256 or not hmac.compare_digest(supplied,expected):
                    raise APIError(403,'Доступ запрещён')
                with connect() as db:
                    signature=backup_tools.fingerprint(db,backup_tools.CONTENT_TABLES)
                    previous=self.headers.get('X-Forma-If-Unchanged','').strip()
                    if len(previous)==64 and hmac.compare_digest(previous,signature):
                        self.send_response(204)
                        self.send_header('Cache-Control','no-store')
                        self.send_header('X-Forma-Data-Signature',signature)
                        self.end_headers()
                        return
                    raw,_manifest=self.make_backup(db)
                name='forma-backup-'+dt.datetime.now(TZ).strftime('%Y%m%d-%H%M%S')+'.zip'
                self.send_response(200)
                self.send_header('Content-Type','application/zip')
                self.send_header('Content-Disposition',f'attachment; filename="{name}"')
                self.send_header('Content-Length',str(len(raw)))
                self.send_header('Cache-Control','no-store')
                self.send_header('X-Content-Type-Options','nosniff')
                self.send_header('X-Forma-Data-Signature',signature)
                self.end_headers()
                return self.wfile.write(raw)
            if not path.startswith('/api/'):
                if method!='GET': raise APIError(405,'Метод не поддерживается')
                return self.serve_static(path)
            with connect() as db:
                if path=='/api/login' and method=='POST': return self.login(db)
                user=self.require(db, path in ('/api/users','/api/tasks','/api/logs','/api/report','/api/export') and path not in ('/api/tasks',)) if path!='/api/session' else self.user(db)
                if path=='/api/session': return self.send_json({'user':dict(user) if user else None,'today':TODAY().isoformat()})
                if path=='/api/logout' and method=='POST':
                    token=self.auth_token()
                    if token:
                        token_hash=hashlib.sha256(token.encode()).hexdigest()
                        db.execute('DELETE FROM sessions WHERE token_hash=?',(token_hash,))
                        db.execute('DELETE FROM prank_presence WHERE session_hash=?',(token_hash,))
                    log(db,user['id'],'logout','Вышел из системы'); db.commit()
                    return self.send_json({'ok':True},cookie='session=; HttpOnly; SameSite=Lax; Path=/; Max-Age=0')
                if path=='/api/sotipb/template' and method=='GET':
                    self.require(db,True)
                    return self.send_bytes(sotipb.template(),'sotipb-template.xlsx','application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
                if path=='/api/sotipb/history' and method=='GET': return self.sotipb_history(db,self.require(db,True),q)
                if path=='/api/sotipb/convert' and method=='POST': return self.sotipb_convert(db,self.require(db,True))
                if path=='/api/sotipb/file' and method=='GET': return self.sotipb_file(db,self.require(db,True),q)
                if path.startswith('/api/sotipb/history/') and method=='DELETE': return self.sotipb_delete(db,self.require(db,True),path)
                if path=='/api/deploy/check' and method=='GET': return self.deploy_check(db,self.require(db,True))
                if path=='/api/backup/status' and method=='GET': return self.backup_status(db,self.require(db,True))
                if path=='/api/backup/download' and method=='GET': return self.backup_download(db,self.require(db,True))
                if path=='/api/backup/safety' and method=='GET': return self.backup_safety(db,self.require(db,True),q)
                if path=='/api/backup/inspect' and method=='POST': return self.backup_inspect(db,self.require(db,True))
                if path=='/api/backup/restore' and method=='POST': return self.backup_restore(db,self.require(db,True),q)
                if path=='/api/prank/poll' and method=='GET': return self.prank_poll(db,user,q)
                if path=='/api/presence' and method=='GET': return self.presence(db,user)
                if path=='/api/prank/online' and method=='GET': return self.prank_online(db,user)
                if path=='/api/prank/send' and method=='POST': return self.prank_send(db,user,self.body())
                if path=='/api/vacation' and method=='GET': return self.vacation_list(db,user,q)
                if path=='/api/vacation/save' and method=='POST': return self.vacation_save(db,self.require(db,True),self.body())
                if path=='/api/vacation/help' and method=='POST': return self.vacation_help(db,self.require(db,True),self.body())
                if path=='/api/vacation/import' and method=='POST': return self.vacation_import(db,self.require(db,True),q)
                if path=='/api/vacation/export' and method=='GET': return self.vacation_export(db,user,q)
                if path=='/api/bootstrap' and method=='GET': return self.bootstrap(db,user,q)
                if path=='/api/calls' and method=='GET': return self.call_status(db,user,q)
                if path=='/api/calls' and method=='POST': return self.call_start(db,user,self.body())
                if path=='/api/calls/fallback' and method=='POST': return self.call_fallback(db,user,self.body())
                if path=='/api/calls/cancel' and method=='POST': return self.call_cancel(db,user,self.body())
                if path=='/api/calls/join' and method=='POST': return self.call_join(db,user,self.body())
                if path=='/api/calls/heartbeat' and method=='POST': return self.call_heartbeat(db,user,self.body())
                if path=='/api/calls/leave' and method=='POST': return self.call_leave(db,user,self.body())
                if path=='/api/calls/signals' and method=='GET': return self.call_signals(db,user,q)
                if path=='/api/calls/signals' and method=='POST': return self.call_signal_send(db,user,self.body())
                if path=='/api/chat/groups' and method=='POST': return self.chat_group_create(db,user,self.body())
                if path.startswith('/api/chat/groups/') and method=='PATCH': return self.chat_group_edit(db,user,path,self.body())
                if path=='/api/chat' and method=='GET': return self.chat_list(db,user,q)
                if path=='/api/chat' and method=='POST': return self.chat_send(db,user,self.body())
                if path=='/api/chat/status' and method=='GET': return self.chat_status(db,user,q)
                if path=='/api/chat/read' and method=='POST': return self.chat_read(db,user,self.body())
                if path=='/api/chat/pin' and method=='POST': return self.chat_pin(db,user,self.body())
                if path=='/api/chat/theme' and method=='POST': return self.chat_theme_save(db,user,self.body())
                if path=='/api/avatar' and method=='POST': return self.avatar_save(db,user,self.body(150000))
                if path.startswith('/api/avatar/') and method=='GET': return self.avatar_image(db,user,path)
                if path=='/api/entries' and method=='POST': return self.save_entry(db,user,self.body())
                if path=='/api/comments' and method=='POST': return self.save_comment(db,user,self.body())
                if path.startswith('/api/entries/') and method=='PATCH': return self.edit_entry(db,user,path,self.body())
                if path.startswith('/api/entries/') and method=='DELETE': return self.delete_entry(db,user,path)
                if path=='/api/attendance' and method=='POST': return self.save_attendance(db,self.require(db,True),self.body())
                if path=='/api/hours' and method=='POST': return self.save_hours(db,self.require(db,True),self.body())
                if path=='/api/users' and method=='POST': return self.add_user(db,self.require(db,True),self.body())
                if path.startswith('/api/users/') and method=='PATCH': return self.edit_user(db,self.require(db,True),path,self.body())
                if path=='/api/tasks' and method=='POST': return self.add_task(db,self.require(db,True),self.body())
                if path.startswith('/api/tasks/') and method=='PATCH': return self.edit_task(db,self.require(db,True),path,self.body())
                if path=='/api/password' and method=='POST': return self.password(db,self.require(db,True),self.body())
                if path=='/api/logs' and method=='GET':
                    self.require(db,True)
                    rows=db.execute('''SELECT a.id,a.action,a.detail,a.created_at,COALESCE(u.name,'Система') actor
                      FROM activity a LEFT JOIN users u ON u.id=a.actor_id ORDER BY a.id DESC LIMIT 150''').fetchall()
                    return self.send_json({'logs':json_rows(rows)})
                if path=='/api/report' and method=='GET': return self.report(db,self.require(db,True),q)
                if path=='/api/export' and method=='GET': return self.export(db,self.require(db,True),q)
                raise APIError(404,'Адрес не найден')
        except APIError as exc: self.send_json({'error':exc.message},exc.status)
        except d1_store.RemoteError as exc: self.send_json({'error':str(exc)},503)
        except sqlite3.IntegrityError: self.send_json({'error':'Конфликт данных. Обновите страницу и повторите действие'},409)
        except (BrokenPipeError,ConnectionResetError): pass
        except Exception:
            traceback.print_exc(); self.send_json({'error':'Ошибка сервера'},500)
    def login(self,db):
        data=self.body(); username=str(data.get('username','')).strip(); password=str(data.get('password',''))
        row=db.execute('SELECT * FROM users WHERE username=? AND active=1',(username,)).fetchone()
        if not row or not valid_password(row['password_hash'],password): raise APIError(401,'Неверный логин или пароль')
        token=secrets.token_urlsafe(32)
        expires=(dt.datetime.now(dt.timezone.utc)+dt.timedelta(days=14)).isoformat()
        db.execute('INSERT INTO sessions(token_hash,user_id,expires_at) VALUES(?,?,?)',(hashlib.sha256(token.encode()).hexdigest(),row['id'],expires))
        log(db,row['id'],'login','Вход в систему'); db.commit()
        self.send_json({'user':{k:row[k] for k in ('id','name','username','role','is_staff','active')},'session_token':token},cookie=f'session={token}; HttpOnly; SameSite=Lax; Path=/; Max-Age=1209600')
    def vacation_year(self,q):
        try: return vacation.valid_year(q.get('year',[str(TODAY().year)])[0])
        except vacation.VacationError as exc: raise APIError(400,str(exc))
    def vacation_list(self,db,user,q):
        year=self.vacation_year(q)
        people,rows=vacation.roster(db,year)
        self.send_json({'year':year,'people':people,'entries':rows})
    def vacation_save(self,db,user,data):
        try:
            year=vacation.valid_year(data.get('year'))
            uid=data.get('user_id');month=data.get('month')
            if type(uid)!=int or type(month)!=int or not 1<=month<=12: raise vacation.VacationError('Некорректный сотрудник или месяц')
            person=db.execute('SELECT name FROM users WHERE id=? AND is_staff=1',(uid,)).fetchone()
            if not person: raise vacation.VacationError('Сотрудник не найден')
            check_expected(db,data,'vacation_schedule','user_id=? AND year=? AND month=?',
                           (uid,year,month),('start_date','end_date','material_help'))
            start=data.get('start_date');end=data.get('end_date')
            if start=='' and end=='':
                db.execute('DELETE FROM vacation_schedule WHERE user_id=? AND year=? AND month=?',(uid,year,month))
                detail=f'Удалён отпуск: {person["name"]}, {year}, {month}'
            else:
                vacation.dates(year,month,start,end)
                db.execute('''INSERT INTO vacation_schedule(user_id,year,month,start_date,end_date)
                    VALUES(?,?,?,?,?) ON CONFLICT(user_id,year,month) DO UPDATE SET
                    start_date=excluded.start_date,end_date=excluded.end_date,updated_at=CURRENT_TIMESTAMP''',
                    (uid,year,month,start,end))
                detail=f'Отпуск: {person["name"]}, {start}—{end}'
            log(db,user['id'],'vacation',detail);db.commit()
        except vacation.VacationError as exc:raise APIError(400,str(exc))
        self.send_json({'ok':True})
    def vacation_help(self,db,user,data):
        try:
            year=vacation.valid_year(data.get('year'))
            uid=data.get('user_id');month=data.get('month');help_flag=data.get('material_help')
            if type(uid)!=int or type(month)!=int or type(help_flag)!=bool or not 1<=month<=12:
                raise vacation.VacationError('Некорректные данные')
            exists=db.execute('''SELECT u.name FROM vacation_schedule v JOIN users u ON u.id=v.user_id
                WHERE v.user_id=? AND v.year=? AND v.month=? AND u.is_staff=1''',(uid,year,month)).fetchone()
            if not exists:raise vacation.VacationError('Сначала укажите даты отпуска в верхней таблице')
            check_expected(db,data,'vacation_schedule','user_id=? AND year=? AND month=?',
                           (uid,year,month),('start_date','end_date','material_help'))
            db.execute('''UPDATE vacation_schedule SET material_help=?,updated_at=CURRENT_TIMESTAMP
                WHERE user_id=? AND year=? AND month=?''',(int(help_flag),uid,year,month))
            log(db,user['id'],'vacation_help',f'{exists["name"]}, {year}-{month:02d}: {"+М" if help_flag else "+"}')
            db.commit()
        except vacation.VacationError as exc:raise APIError(400,str(exc))
        self.send_json({'ok':True})
    def vacation_import(self,db,user,q):
        year=self.vacation_year(q)
        if self.headers.get('Content-Type','').split(';')[0]!='application/octet-stream':
            raise APIError(415,'Загрузите XLSX')
        try:size=int(self.headers.get('Content-Length','0'))
        except ValueError:raise APIError(400,'Некорректный размер XLSX')
        if not 0<size<=2*1024*1024:raise APIError(413,'Размер XLSX не более 2 МБ')
        raw=self.rfile.read(size)
        if len(raw)!=size:raise APIError(400,'Файл не загрузился полностью')
        try:
            people,_=vacation.roster(db,year)
            map_arg=q.get('mapping',['{}'])[0]
            if len(map_arg)>4000:raise vacation.VacationError('Слишком длинное сопоставление ФИО')
            overrides=json.loads(map_arg)
            if not isinstance(overrides,dict):raise vacation.VacationError('Некорректное сопоставление ФИО')
            result=vacation.parse_xlsx(raw,people,year,overrides)
        except (vacation.VacationError,ValueError,TypeError) as exc:raise APIError(400,str(exc))
        existing={(r['user_id'],r['month']) for r in db.execute('SELECT user_id,month FROM vacation_schedule WHERE year=?',(year,))}
        snapshot_rows=[dict(r) for r in db.execute('''SELECT user_id,month,start_date,end_date,material_help
              FROM vacation_schedule WHERE year=? ORDER BY user_id,month''',(year,))]
        snapshot=hashlib.sha256(json.dumps(snapshot_rows,ensure_ascii=False,sort_keys=True).encode()).hexdigest()
        if q.get('preview',['1'])[0]=='0' and 'snapshot' in q and q['snapshot'][0]!=snapshot:
            raise APIError(409,'План отпусков изменился после предпросмотра. Обновите предпросмотр и проверьте его перед импортом')
        pending=[r for r in result['entries'] if (r[0],r[2]) not in existing]
        payload={'year':year,'matched':result['matched'],'unmatched':result['unmatched'],
                 'found':len(result['entries']),'mapping':result['mapping'],'snapshot':snapshot,'new':len(pending),'unchanged':len(result['entries'])-len(pending)}
        if q.get('preview',['1'])[0]=='0':
            for row in pending:
                db.execute('''INSERT OR IGNORE INTO vacation_schedule
                  (user_id,year,month,start_date,end_date,material_help) VALUES(?,?,?,?,?,?)''',row)
            log(db,user['id'],'vacation_import',f'{year}: добавлено {len(pending)}, пропущено ранее заполненных {payload["unchanged"]}')
            db.commit()
        return self.send_json(payload)
    def vacation_export(self,db,user,q):
        year=self.vacation_year(q)
        people,rows=vacation.roster(db,year)
        raw=vacation.export_xlsx(year,people,rows)
        self.send_bytes(raw,f'forma-otpusk-{year}.xlsx',
                        'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    def sotipb_history(self,db,user,q):
        try:before=int(q.get('before',['0'])[0])
        except ValueError:raise APIError(400,'Некорректная страница истории')
        if before<0:raise APIError(400,'Некорректная страница истории')
        runs=json_rows(db.execute('''SELECT r.id,r.source_name,r.source_size,r.result_size,
          r.row_count,r.group_count,r.warnings_json,r.created_at,u.name uploader_name
          FROM sotipb_runs r LEFT JOIN users u ON u.id=r.uploader_id
          WHERE r.state='ready' AND (?=0 OR r.id<?) ORDER BY r.id DESC LIMIT 51''',
          (before,before)).fetchall())
        more=len(runs)>50;runs=runs[:50]
        for run in runs:run['warnings']=json.loads(run.pop('warnings_json'))
        return self.send_json({'runs':runs,'has_more':more,
                               'next':runs[-1]['id'] if more else None,'limit_bytes':12*1024*1024})
    def sotipb_convert(self,db,user):
        if self.headers.get('Content-Type','').split(';')[0]!='application/octet-stream':
            raise APIError(415,'Выберите XLSX-файл')
        try:size=int(self.headers.get('Content-Length','0'))
        except ValueError:raise APIError(400,'Некорректный размер файла')
        if not 0<size<=sotipb.MAX_INPUT:raise APIError(413,'Исходный XLSX: максимум 3 МБ')
        source=self.rfile.read(size)
        if len(source)!=size:raise APIError(400,'Файл загрузился не полностью')
        filename=self.headers.get('X-Forma-File-Name','')
        from urllib.parse import unquote
        filename=re.sub(r'[\\/\r\n\x00-\x1f]','',unquote(filename)).strip()[:140]
        if not filename.lower().endswith('.xlsx'):filename='лист 1.xlsx'
        try:result,details=sotipb.convert(source)
        except sotipb.ConversionError as exc:raise APIError(400,str(exc))
        cap=12*1024*1024
        used=db.execute("SELECT COALESCE(SUM(source_size+result_size),0) FROM sotipb_runs WHERE state='ready'").fetchone()[0]
        if used+len(source)+len(result)>cap:
            raise APIError(413,'История СОТИПБ достигла 12 МБ. Скачайте нужные файлы и удалите старые пары перед новой загрузкой')
        args=(user['id'],filename,len(source),len(result),details['rows'],details['groups'],
              json.dumps(details['warnings'],ensure_ascii=False),hashlib.sha256(source).hexdigest(),hashlib.sha256(result).hexdigest())
        insert='''INSERT INTO sotipb_runs(uploader_id,source_name,source_size,result_size,
                row_count,group_count,warnings_json,source_sha256,result_sha256,state)
                VALUES(?,?,?,?,?,?,?,?,?,'pending')'''
        run=None
        try:
            cursor=db.query(insert,args) if isinstance(db,d1_store.D1Connection) else db.execute(insert,args)
            run=cursor.lastrowid
            for kind,content in (('source',source),('result',result)):
                for part,start in enumerate(range(0,len(content),48000)):
                    db.execute('INSERT INTO sotipb_parts(run_id,kind,part_no,content_b64) VALUES(?,?,?,?)',
                               (run,kind,part,base64.b64encode(content[start:start+48000]).decode('ascii')))
            db.commit()
            db.execute("UPDATE sotipb_runs SET state='ready' WHERE id=?",(run,))
            log(db,user['id'],'sotipb',f'СОТИПБ: {filename}, {details["rows"]} строк, {details["groups"]} групп')
            db.commit()
        except Exception:
            if run is not None:
                try:
                    db.rollback();db.execute('DELETE FROM sotipb_parts WHERE run_id=?',(run,))
                    db.execute('DELETE FROM sotipb_runs WHERE id=?',(run,));db.commit()
                except Exception:pass
            raise
        return self.send_json({'ok':True,'id':run,'rows':details['rows'],
                               'groups':details['groups'],'warnings':details['warnings'],
                               'warning_count':details['warning_count']})
    def sotipb_file(self,db,user,q):
        try:run=int(q.get('id',[''])[0]);kind=q.get('kind',[''])[0]
        except ValueError:raise APIError(400,'Некорректный номер файла')
        if run<=0 or kind not in ('source','result'):raise APIError(400,'Некорректный номер файла')
        row=db.execute('SELECT * FROM sotipb_runs WHERE id=? AND state=?',(run,'ready')).fetchone()
        if not row:raise APIError(404,'Запись истории не найдена')
        chunks=db.execute('SELECT part_no,content_b64 FROM sotipb_parts WHERE run_id=? AND kind=? ORDER BY part_no',(run,kind)).fetchall()
        try:content=b''.join(base64.b64decode(part['content_b64'],validate=True) for part in chunks)
        except (ValueError,binascii.Error):raise APIError(503,'Сохранённый файл повреждён')
        if [part['part_no'] for part in chunks]!=list(range(len(chunks))) or len(content)!=row[kind+'_size'] or hashlib.sha256(content).hexdigest()!=row[kind+'_sha256']:
            raise APIError(503,'Сохранённый файл не прошёл проверку целостности')
        return self.send_bytes(content,f'sotipb-{kind}-{run}.xlsx',
                               'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    def sotipb_delete(self,db,user,path):
        try:run=int(path.rsplit('/',1)[1])
        except ValueError:raise APIError(400,'Некорректный номер файла')
        if run<=0:raise APIError(400,'Некорректный номер файла')
        row=db.execute('SELECT source_name FROM sotipb_runs WHERE id=? AND state=?',(run,'ready')).fetchone()
        if not row:raise APIError(404,'Запись истории не найдена')
        db.execute('DELETE FROM sotipb_parts WHERE run_id=?',(run,))
        db.execute('DELETE FROM sotipb_runs WHERE id=?',(run,))
        log(db,user['id'],'sotipb_delete',f'Удалена пара файлов СОТИПБ #{run}: {row["source_name"]}')
        db.commit();return self.send_json({'ok':True})
    def deploy_check(self,db,user):
        """Read-only post-deploy diagnostics on the real running process and its D1."""
        required=('/index.html','/main.js','/calls.js','/livekit-client.umd.min.js','/styles.css','/effect-snake.png')
        missing=[name for name in required if not (ROOT/'web'/name[1:]).is_file()]
        if not sotipb.SUMMARY_LAYOUT.is_file():missing.append('sotipb-summary-layout.xlsx')
        index=(ROOT/'web'/'index.html').read_text(encoding='utf-8')
        try:
            if isinstance(db,d1_store.D1Connection):d1_store.require_complete_schema(db)
            counts=backup_tools.summary(db)['counts']
            database=True;detail='D1 отвечает; таблицы отпуска и чата доступны'
        except Exception as exc:
            counts={};database=False;detail='База или набор таблиц недоступны: '+str(exc)[:170]
        checks={'render':True,'interface':not missing and ('main.js?v='+BUILD_ID) in index and ('styles.css?v='+BUILD_ID) in index,
                'database':database,'backup_tables':database and set(backup_tools.BUSINESS)==set(d1_store.TABLES)-d1_store.EXCLUDED_FROM_BACKUP}
        self.send_json({'build':BUILD_ID,'backend':os.environ.get('DATABASE_BACKEND','sqlite'),
                        'checks':checks,'ok':all(checks.values()),'detail':detail,
                        'missing_files':missing,'counts':{'vacation_schedule':counts.get('vacation_schedule'),
                                                        'chat_messages':counts.get('chat_messages')}})
    def backup_status(self,db,user):
        safety=[]
        if BACKUP_DIR.is_dir():
            for item in sorted(BACKUP_DIR.glob('before-restore-*.zip'),reverse=True)[:20]:
                if item.is_file() and not item.is_symlink():
                    safety.append({'name':item.name,'size':item.stat().st_size})
        monitor=json_rows(db.execute('SELECT source,last_check,last_saved,file_name,file_size,signature,state,detail FROM backup_monitor ORDER BY source').fetchall())
        self.send_json({'current':backup_tools.summary(db),'safety':safety,'monitor':monitor})
    def backup_safety(self,db,user,q):
        name=q.get('name',[''])[0]
        if not re.fullmatch(r'before-restore-\d{8}-\d{6}-[0-9a-f]{8}\.zip',name):
            raise APIError(400,'Некорректное имя резервной копии')
        file=BACKUP_DIR/name
        if not file.is_file() or file.is_symlink(): raise APIError(404,'Копия не найдена')
        self.send_bytes(file.read_bytes(),name,'application/zip')
    def make_backup(self,db):
        if isinstance(db,d1_store.D1Connection):
            snapshot=d1_store.sqlite_snapshot(db)
            try: return backup_tools.archive(snapshot)
            finally: snapshot.close()
        return backup_tools.archive(db)
    def backup_download(self,db,user):
        raw,manifest=self.make_backup(db)
        name='forma-backup-'+dt.datetime.now(TZ).strftime('%Y%m%d-%H%M%S')+'.zip'
        db.execute('''INSERT INTO backup_monitor(source,last_check,last_saved,file_name,file_size,state)
            VALUES('manual',CURRENT_TIMESTAMP,CURRENT_TIMESTAMP,?,?,'generated')
            ON CONFLICT(source) DO UPDATE SET last_check=CURRENT_TIMESTAMP,last_saved=CURRENT_TIMESTAMP,
            file_name=excluded.file_name,file_size=excluded.file_size,state='generated'
            ''',(name,len(raw)))
        db.commit()
        self.send_bytes(raw,name,'application/zip')
    def backup_inspect(self,db,user):
        raw=self.backup_body()
        try: _data,manifest,saved=backup_tools.inspect(raw)
        except backup_tools.BackupError as exc: raise APIError(400,str(exc))
        current=backup_tools.summary(db)
        self.send_json({'created_at':manifest['created_at'],'backup':saved,'current':current,
                        'comparison':backup_tools.compare(current,saved),
                        'archive_hash':hashlib.sha256(raw).hexdigest(),
                        'current_signature':backup_tools.fingerprint(db)})
    def backup_restore(self,db,user,q):
        if self.headers.get('X-Forma-Confirm')!='RESTORE':
            raise APIError(400,'Подтвердите полную замену данных')
        raw=self.backup_body()
        try: data,manifest,saved=backup_tools.inspect(raw)
        except backup_tools.BackupError as exc: raise APIError(400,str(exc))
        supplied_hash=q.get('archive_hash',[''])[0]
        supplied_signature=q.get('current_signature',[''])[0]
        if not hmac.compare_digest(hashlib.sha256(raw).hexdigest(),supplied_hash):
            raise APIError(400,'Архив изменился после сравнения. Сравните данные ещё раз')
        if not hmac.compare_digest(backup_tools.fingerprint(db),supplied_signature):
            raise APIError(409,'Данные сайта изменились после сравнения. Сравните архив ещё раз')
        # Страховочная копия на сервере: скачайте свою копию отдельно.
        safety,_=self.make_backup(db)
        BACKUP_DIR.mkdir(mode=0o700,parents=True,exist_ok=True)
        safe_name='before-restore-'+dt.datetime.now(TZ).strftime('%Y%m%d-%H%M%S')+'-'+secrets.token_hex(4)+'.zip'
        safe_path=BACKUP_DIR/safe_name
        with safe_path.open('xb') as output: output.write(safety)
        safe_path.chmod(0o600)
        with sqlite3.connect(':memory:') as source:
            source.deserialize(data)
            backup_tools.validate(source)  # Дополняет старые архивы пустой таблицей чата.
            # Old external ZIPs may contain avatars; never restore them into current data.
            source.execute('DELETE FROM user_avatars')
            source.commit()
            expected_signature=backup_tools.fingerprint(source)
            if isinstance(db,d1_store.D1Connection): d1_store.replace_atomic(db,source)
            else: source.backup(db)
        # Verify all persistent tables, including vacations, chart inputs and chat,
        # before reporting a restore as successful. Activity is logged just below.
        if backup_tools.fingerprint(db)!=expected_signature:
            raise APIError(503,'Данные восстановлены, но полная проверка не прошла. Не загружайте ZIP повторно; обратитесь за помощью')
        if not isinstance(db,d1_store.D1Connection):
            db.execute('DELETE FROM sessions')
            db.execute('DELETE FROM prank_presence')
            db.execute('DELETE FROM prank_events')
        log(db,None,'backup_restored',f'Восстановлено из копии {manifest["created_at"]}; перед заменой: {safe_name}')
        db.commit()
        return self.send_json({'ok':True,'safety_copy':safe_name,'restored':saved})
    def prank_poll(self,db,user,q):
        """Heartbeat and short-lived event stream for open browser tabs."""
        now=dt.datetime.now(dt.timezone.utc)
        token_hash=hashlib.sha256(self.auth_token().encode()).hexdigest()
        initial='after' not in q
        try:
            after=int(q.get('after',['0'])[0])
            if after<0: raise ValueError()
        except (ValueError,TypeError): raise APIError(400,'Некорректный номер события')
        db.execute('''INSERT INTO prank_presence(session_hash,user_id,last_seen) VALUES(?,?,?)
          ON CONFLICT(session_hash) DO UPDATE SET user_id=excluded.user_id,last_seen=excluded.last_seen''',
          (token_hash,user['id'],now.isoformat()))
        latest=db.execute('SELECT COALESCE(MAX(id),0) FROM prank_events').fetchone()[0]
        # Первая проверка вкладки подписывает только на будущие события.
        events=[] if initial else json_rows(db.execute('''SELECT id,kind,body FROM prank_events
          WHERE target_id=? AND id>? AND created_at>? ORDER BY id LIMIT 20''',
          (user['id'],after,(now-dt.timedelta(minutes=2)).isoformat())).fetchall())
        for event in events:
            if event['kind']=='speech' and event['body'].startswith(PRANK_EFFECT_PREFIX):
                try:
                    packed=json.loads(event['body'][len(PRANK_EFFECT_PREFIX):])
                    if packed.get('kind') in PRANK_NEW_KINDS and isinstance(packed.get('text'),str):
                        event['kind']=packed['kind'];event['body']=packed['text']
                except (ValueError,TypeError,AttributeError): pass
        db.commit()
        return self.send_json({'events':events,'last_id':max(latest,after)})
    def presence(self,db,user):
        """Authenticated roster: only active sessions with a recent browser heartbeat."""
        now=dt.datetime.now(dt.timezone.utc)
        rows=db.execute('''SELECT DISTINCT u.id,u.name,u.role,
          COALESCE(a.revision,0) avatar_revision,COALESCE(a.image_url,'') avatar_url,
          CASE WHEN a.image_b64<>'' THEN 1 ELSE 0 END avatar_stored
          FROM users u JOIN prank_presence p ON p.user_id=u.id
          JOIN sessions s ON s.token_hash=p.session_hash
          LEFT JOIN user_avatars a ON a.user_id=u.id
          WHERE u.active=1 AND p.last_seen>=? AND s.expires_at>?
          ORDER BY CASE WHEN u.role='admin' THEN 0 ELSE 1 END,u.name''',
          ((now-PRESENCE_WINDOW).isoformat(),now.isoformat())).fetchall()
        return self.send_json({'users':json_rows(rows)})
    def prank_online(self,db,user):
        """Directory of active targets for logged-in users; includes online status."""
        now=dt.datetime.now(dt.timezone.utc)
        cutoff=(now-PRESENCE_WINDOW).isoformat()
        rows=json_rows(db.execute('''SELECT u.id,u.name,u.role,
            CASE WHEN EXISTS(SELECT 1 FROM prank_presence p
               JOIN sessions s ON s.token_hash=p.session_hash
               WHERE p.user_id=u.id AND p.last_seen>=? AND s.expires_at>?) THEN 1 ELSE 0 END online
          FROM users u WHERE u.active=1 ORDER BY u.name,u.id''',
          (cutoff,now.isoformat())).fetchall())
        return self.send_json({'user_ids':[r['id'] for r in rows if r['online']],
                               'users':[{'id':r['id'],'name':r['name'],'role':r['role']} for r in rows]})
    def prank_send(self,db,user,data):
        try: target=int(data.get('user_id'))
        except (TypeError,ValueError): raise APIError(400,'Выберите сотрудника')
        if isinstance(data.get('user_id'),bool): raise APIError(400,'Выберите сотрудника')
        kind=data.get('kind')
        if kind not in ('people','speech',*PRANK_NEW_KINDS): raise APIError(400,'Неизвестный эффект')
        body=data.get('text','')
        if not isinstance(body,str): raise APIError(400,'Текст должен быть строкой')
        body=body.strip()
        if kind=='speech' and (not body or len(body)>400): raise APIError(400,'Введите текст от 1 до 400 символов')
        if kind=='cat' and (not body or len(body)>140): raise APIError(400,'Сообщение кота: от 1 до 140 символов')
        if kind=='achievement' and (not body or len(body)>80): raise APIError(400,'Название достижения: от 1 до 80 символов')
        if kind=='plane' and (not body or len(body)>140): raise APIError(400,'Послание самолётика: от 1 до 140 символов')
        if kind=='stamp' and (not body or len(body)>65): raise APIError(400,'Текст печати: от 1 до 65 символов')
        if kind=='duck' and (not body or len(body)>100): raise APIError(400,'Реплика утки: от 1 до 100 символов')
        if kind=='bonus' and (not body or len(body)>120): raise APIError(400,'Секретный бонус: от 1 до 120 символов')
        if kind=='stickers' and body not in ('stars','hearts'): raise APIError(400,'Выберите звёзды или сердечки')
        if kind in ('people','parade'): body=''
        stored_kind=kind if kind in ('people','speech') else 'speech'
        stored_body=PRANK_EFFECT_PREFIX+json.dumps({'kind':kind,'text':body},ensure_ascii=False) if kind in PRANK_NEW_KINDS else body
        recipient=db.execute('SELECT name FROM users WHERE id=? AND active=1',(target,)).fetchone()
        if not recipient: raise APIError(404,'Пользователь не найден')
        now=dt.datetime.now(dt.timezone.utc)
        present=db.execute('''SELECT 1 FROM prank_presence p JOIN sessions s ON s.token_hash=p.session_hash
          WHERE p.user_id=? AND p.last_seen>=? AND s.expires_at>? LIMIT 1''',
          (target,(now-PRESENCE_WINDOW).isoformat(),now.isoformat())).fetchone()
        if not present: raise APIError(409,'Пользователь сейчас не в сети. Попросите его открыть сайт.')
        db.execute('INSERT INTO prank_events(target_id,kind,body,created_at) VALUES(?,?,?,?)',
                   (target,stored_kind,stored_body,now.isoformat()))
        db.execute('DELETE FROM prank_events WHERE created_at<?',((now-dt.timedelta(days=1)).isoformat(),))
        label={'people':'Человечки','speech':'Голосовое сообщение','cat':'Кот-курьер',
               'achievement':'Достижение','parade':'Мини-парад','plane':'Бумажный самолётик',
               'stamp':'Печать «Одобрено!»','duck':'Важная утка','stickers':'Дождь из стикеров',
               'bonus':'Секретный бонус'}[kind]
        log(db,user['id'],'prank',f'{label} → {recipient["name"]}')
        db.commit()
        return self.send_json({'ok':True,'name':recipient['name']})
    def bootstrap(self,db,user,q):
        month=q.get('month',[TODAY().strftime('%Y-%m')])[0]; first,last=month_range(month)
        admin=user['role']=='admin'
        users=json_rows(db.execute('''SELECT u.id,u.name,u.username,u.role,u.is_staff,u.active,u.created_at,u.password_hash,
            COALESCE(a.revision,0) avatar_revision,COALESCE(a.image_url,'') avatar_url,
            CASE WHEN a.image_b64<>'' THEN 1 ELSE 0 END avatar_stored
            FROM users u LEFT JOIN user_avatars a ON a.user_id=u.id
            WHERE u.is_staff=1 ORDER BY u.active DESC,u.name''').fetchall()) if admin else [dict(user)]
        if admin:
            for visible in users:visible['password_revision']=hashlib.sha256(visible.pop('password_hash').encode()).hexdigest()
        tasks=json_rows(db.execute('SELECT id,title,unit,norm,category,active FROM tasks ORDER BY active DESC,id').fetchall())
        eargs=[first,last]; aargs=[first,last]
        efilter=''; afilter=''
        if not admin:
            efilter=' AND e.user_id=?'; afilter=' AND a.user_id=?'; eargs.append(user['id']); aargs.append(user['id'])
        entries=json_rows(db.execute('''SELECT e.id,e.user_id,e.task_id,e.work_date,e.quantity,e.note,e.created_at,e.updated_at
            FROM entries e WHERE e.work_date BETWEEN ? AND ?'''+efilter+' ORDER BY e.work_date DESC,e.id DESC',eargs).fetchall())
        comment_args=[first,last] if admin else [first,last,user['id']]
        comments=json_rows(db.execute('''SELECT user_id,task_id,work_date,body FROM cell_comments
            WHERE work_date BETWEEN ? AND ?'''+('' if admin else ' AND user_id=?'),comment_args).fetchall())
        attendance=json_rows(db.execute('SELECT a.user_id,a.day,a.status,a.reason FROM attendance a WHERE a.day BETWEEN ? AND ?'+afilter,aargs).fetchall())
        daily_hours=json_rows(db.execute('SELECT day,hours FROM daily_hours WHERE day BETWEEN ? AND ?',(first,last)).fetchall())
        hours_filter='' if admin else ' AND user_id=?'
        hours_args=[first,last] if admin else [first,last,user['id']]
        personal_hours=json_rows(db.execute('SELECT user_id,day,hours FROM personal_hours WHERE day BETWEEN ? AND ?'+hours_filter,hours_args).fetchall())
        recent=json_rows(db.execute('''SELECT a.action,a.detail,a.created_at,COALESCE(u.name,'Система') actor
            FROM activity a LEFT JOIN users u ON u.id=a.actor_id ORDER BY a.id DESC LIMIT 6''').fetchall()) if admin else []
        return self.send_json({'user':dict(user),'users':users,'tasks':tasks,'entries':entries,'comments':comments,'attendance':attendance,
            'daily_hours':daily_hours,'personal_hours':personal_hours,'recent':recent,'month':month,'today':TODAY().isoformat()})
    def save_entry(self,db,user,data):
        task_id=int(data.get('task_id') or 0); day=date_ok(data.get('date'))
        try: qty=float(data.get('quantity'))
        except (ValueError,TypeError): raise APIError(400,'Укажите количество')
        if not 0<qty<=10000000 or not __import__('math').isfinite(qty): raise APIError(400,'Количество должно быть больше нуля')
        note=str(data.get('note','')).strip()[:500]
        task=db.execute('SELECT id,title,active FROM tasks WHERE id=?',(task_id,)).fetchone()
        if not task or not task['active']: raise APIError(400,'Работа не найдена или архивирована')
        target=int(data.get('user_id') or user['id'])
        if target!=user['id'] and user['role']!='admin': raise APIError(403,'Можно изменять только свои записи')
        owner=db.execute('SELECT name FROM users WHERE id=? AND is_staff=1 AND active=1',(target,)).fetchone()
        if not owner: raise APIError(400,'Сотрудник не найден')
        check_expected(db,data,'entries','user_id=? AND task_id=? AND work_date=?',
                       (target,task_id,day.isoformat()),('quantity','note'))
        db.execute('''INSERT INTO entries(user_id,task_id,work_date,quantity,note) VALUES(?,?,?,?,?)
            ON CONFLICT(user_id,task_id,work_date) DO UPDATE SET quantity=excluded.quantity,note=excluded.note,updated_at=CURRENT_TIMESTAMP''',
            (target,task_id,day.isoformat(),qty,note))
        if note:
            db.execute('''INSERT INTO cell_comments(user_id,task_id,work_date,body) VALUES(?,?,?,?)
                ON CONFLICT(user_id,task_id,work_date) DO UPDATE SET body=excluded.body,updated_at=CURRENT_TIMESTAMP''',
                (target,task_id,day.isoformat(),note))
        elif 'note' in data:
            db.execute('DELETE FROM cell_comments WHERE user_id=? AND task_id=? AND work_date=?',(target,task_id,day.isoformat()))
        log(db,user['id'],'entry',f'{owner["name"]}: {task["title"][:85]} — {qty:g} ({day.isoformat()})'); db.commit()
        return self.send_json({'ok':True})
    def save_comment(self,db,user,data):
        day=date_ok(data.get('date')).isoformat()
        try: task_id=int(data.get('task_id')); target=int(data.get('user_id') or user['id'])
        except (ValueError,TypeError): raise APIError(400,'Некорректная работа или сотрудник')
        if target!=user['id'] and user['role']!='admin': raise APIError(403,'Можно комментировать только свою таблицу')
        owner=db.execute('SELECT name FROM users WHERE id=? AND is_staff=1 AND active=1',(target,)).fetchone()
        task=db.execute('SELECT title,active FROM tasks WHERE id=?',(task_id,)).fetchone()
        if not owner or not task: raise APIError(400,'Сотрудник или работа не найдены')
        if not task['active'] and not db.execute('''SELECT 1 FROM entries WHERE user_id=? AND task_id=? AND work_date=?
            UNION ALL SELECT 1 FROM cell_comments WHERE user_id=? AND task_id=? AND work_date=?''',
            (target,task_id,day,target,task_id,day)).fetchone():
            raise APIError(400,'Архивная работа недоступна для новой записи')
        comment=str(data.get('comment') or '').strip()
        if len(comment)>500: raise APIError(400,'Комментарий не длиннее 500 символов')
        check_expected(db,data,'cell_comments','user_id=? AND task_id=? AND work_date=?',
                       (target,task_id,day),('body',))
        if comment:
            db.execute('''INSERT INTO cell_comments(user_id,task_id,work_date,body) VALUES(?,?,?,?)
                ON CONFLICT(user_id,task_id,work_date) DO UPDATE SET body=excluded.body,updated_at=CURRENT_TIMESTAMP''',
                (target,task_id,day,comment))
        else:
            db.execute('DELETE FROM cell_comments WHERE user_id=? AND task_id=? AND work_date=?',(target,task_id,day))
        db.execute('UPDATE entries SET note=?,updated_at=CURRENT_TIMESTAMP WHERE user_id=? AND task_id=? AND work_date=?',
            (comment,target,task_id,day))
        log(db,user['id'],'comment',f'{owner["name"]}: {task["title"][:80]} ({day}) — {"комментарий сохранён" if comment else "комментарий удалён"}')
        db.commit(); self.send_json({'ok':True})
    def edit_entry(self,db,user,path,data):
        try: entry_id=int(path.rsplit('/',1)[1])
        except ValueError: raise APIError(400,'Некорректный идентификатор')
        existing=db.execute('SELECT * FROM entries WHERE id=?',(entry_id,)).fetchone()
        if not existing: raise APIError(404,'Запись не найдена')
        check_expected(db,data,'entries','id=?',(entry_id,),('user_id','task_id','work_date','quantity','note'))
        if existing['user_id']!=user['id'] and user['role']!='admin': raise APIError(403,'Нет доступа')
        day=date_ok(data.get('date'))
        try: task_id=int(data.get('task_id') or 0); target=int(data.get('user_id') or existing['user_id']); qty=float(data.get('quantity'))
        except (ValueError,TypeError): raise APIError(400,'Проверьте заполненные поля')
        if not 0<qty<=10000000 or not __import__('math').isfinite(qty): raise APIError(400,'Количество должно быть больше нуля')
        if target!=existing['user_id'] and user['role']!='admin': raise APIError(403,'Нет доступа')
        task=db.execute('SELECT title,active FROM tasks WHERE id=?',(task_id,)).fetchone()
        owner=db.execute('SELECT name FROM users WHERE id=? AND is_staff=1',(target,)).fetchone()
        if not task or (not task['active'] and task_id!=existing['task_id']) or not owner: raise APIError(400,'Работа или сотрудник не найдены')
        note=str(data.get('note','')).strip()[:500]
        try:
            db.execute('''UPDATE entries SET user_id=?,task_id=?,work_date=?,quantity=?,note=?,updated_at=CURRENT_TIMESTAMP WHERE id=?''',
                (target,task_id,day.isoformat(),qty,note,entry_id))
        except sqlite3.IntegrityError: raise APIError(409,'Запись для этого сотрудника, дня и работы уже существует')
        old_key=(existing['user_id'],existing['task_id'],existing['work_date'])
        new_key=(target,task_id,day.isoformat())
        if old_key!=new_key: db.execute('DELETE FROM cell_comments WHERE user_id=? AND task_id=? AND work_date=?',old_key)
        if note:
            db.execute('''INSERT INTO cell_comments(user_id,task_id,work_date,body) VALUES(?,?,?,?)
                ON CONFLICT(user_id,task_id,work_date) DO UPDATE SET body=excluded.body,updated_at=CURRENT_TIMESTAMP''',(*new_key,note))
        else: db.execute('DELETE FROM cell_comments WHERE user_id=? AND task_id=? AND work_date=?',new_key)
        log(db,user['id'],'entry',f'Исправлено: {owner["name"]}: {task["title"][:85]} — {qty:g} ({day.isoformat()})');db.commit()
        return self.send_json({'ok':True})
    def delete_entry(self,db,user,path):
        try: entry_id=int(path.rsplit('/',1)[1])
        except ValueError: raise APIError(400,'Некорректный идентификатор')
        row=db.execute('SELECT e.*,t.title FROM entries e JOIN tasks t ON t.id=e.task_id WHERE e.id=?',(entry_id,)).fetchone()
        if not row: raise APIError(404,'Запись не найдена')
        if int(self.headers.get('Content-Length','0') or 0):
            check_expected(db,self.body(),'entries','id=?',(entry_id,),('user_id','task_id','work_date','quantity','note'))
        if row['user_id']!=user['id'] and user['role']!='admin': raise APIError(403,'Нет доступа')
        db.execute('DELETE FROM entries WHERE id=?',(entry_id,)); log(db,user['id'],'delete_entry',f'Удалена запись: {row["title"][:85]} ({row["work_date"]})'); db.commit()
        self.send_json({'ok':True})
    def save_attendance(self,db,user,data):
        if user['role']!='admin': raise APIError(403,'Табель изменяет только администратор')
        day=date_ok(data.get('date'),True)
        target=int(data.get('user_id') or user['id'])
        owner=db.execute('SELECT name FROM users WHERE id=? AND is_staff=1 AND active=1',(target,)).fetchone()
        if not owner: raise APIError(400,'Сотрудник не найден')
        status=data.get('status'); reason=str(data.get('reason','')).strip()[:160]
        if status not in ('present','absent','clear'): raise APIError(400,'Некорректный статус')
        if status=='absent' and not reason: raise APIError(400,'Укажите причину отсутствия')
        check_expected(db,data,'attendance','user_id=? AND day=?',(target,day.isoformat()),('status','reason'))
        if status=='clear': db.execute('DELETE FROM attendance WHERE user_id=? AND day=?',(target,day.isoformat()))
        else: db.execute('''INSERT INTO attendance(user_id,day,status,reason) VALUES(?,?,?,?)
            ON CONFLICT(user_id,day) DO UPDATE SET status=excluded.status,reason=excluded.reason,updated_at=CURRENT_TIMESTAMP''',
            (target,day.isoformat(),status,reason if status=='absent' else ''))
        log(db,user['id'],'attendance',f'{owner["name"]}: {day.isoformat()} — {"отсутствует: "+reason if status=="absent" else "на работе" if status=="present" else "по графику"}')
        db.commit(); self.send_json({'ok':True})
    def save_hours(self,db,user,data):
        if 'hours' not in data: raise APIError(400,'Укажите часы или сброс к графику')
        day=date_ok(data.get('date'),True).isoformat()
        target=data.get('user_id')
        if target not in (None,''):
            try: target=int(target)
            except (ValueError,TypeError): raise APIError(400,'Некорректный сотрудник')
            owner=db.execute('SELECT name FROM users WHERE id=? AND is_staff=1',(target,)).fetchone()
            if not owner: raise APIError(400,'Сотрудник не найден')
            label=owner['name']
        else:
            target=None; label='Вся команда'
        if target is None:check_expected(db,data,'daily_hours','day=?',(day,),('hours',))
        else:check_expected(db,data,'personal_hours','user_id=? AND day=?',(target,day),('hours',))
        if data.get('hours') is None:
            if target is None: db.execute('DELETE FROM daily_hours WHERE day=?',(day,))
            else: db.execute('DELETE FROM personal_hours WHERE user_id=? AND day=?',(target,day))
            detail='сброс к графику'
        else:
            try: hours=float(data['hours'])
            except (ValueError,TypeError): raise APIError(400,'Введите количество часов')
            if isinstance(data['hours'],bool) or not math.isfinite(hours) or not 0<=hours<=24 or round(hours,2)!=hours:
                raise APIError(400,'Количество часов: от 0 до 24, не более двух знаков после запятой')
            hours=round(hours,2)
            if target is None:
                db.execute('''INSERT INTO daily_hours(day,hours) VALUES(?,?)
                    ON CONFLICT(day) DO UPDATE SET hours=excluded.hours,updated_at=CURRENT_TIMESTAMP''',(day,hours))
            else:
                db.execute('''INSERT INTO personal_hours(user_id,day,hours) VALUES(?,?,?)
                    ON CONFLICT(user_id,day) DO UPDATE SET hours=excluded.hours,updated_at=CURRENT_TIMESTAMP''',(target,day,hours))
            detail=f'{hours:g} ч'
        log(db,user['id'],'hours',f'{label}: {day} — {detail}')
        db.commit(); self.send_json({'ok':True})
    def add_user(self,db,user,data):
        name=str(data.get('name','')).strip()[:90]; username=str(data.get('username','')).strip().lower()
        password=str(data.get('password',''))
        if len(name)<2 or not USERNAME_PATTERN.fullmatch(username) or len(password)<8: raise APIError(400,'Укажите имя, логин (от 3 символов) и пароль (от 8 символов)')
        try: row=db.execute('INSERT INTO users(name,username,password_hash,role) VALUES(?,?,?,"employee")',(name,username,hashed_password(password)))
        except sqlite3.IntegrityError: raise APIError(409,'Такой логин уже существует')
        log(db,user['id'],'user_added',f'Добавлен сотрудник {name} (@{username})'); db.commit(); self.send_json({'ok':True,'id':row.lastrowid})
    def edit_user(self,db,user,path,data):
        try: uid=int(path.rsplit('/',1)[1])
        except ValueError: raise APIError(400,'Некорректный пользователь')
        target=db.execute('SELECT * FROM users WHERE id=? AND is_staff=1',(uid,)).fetchone()
        if not target: raise APIError(404,'Сотрудник не найден')
        expected=data.get('expected')
        if isinstance(expected,dict) and 'password_revision' in expected:
            revision=hashlib.sha256(target['password_hash'].encode()).hexdigest()
            if expected['password_revision']!=revision:
                raise APIError(409,'Пароль сотрудника уже изменён в другой вкладке. Обновите данные')
            data={**data,'expected':{k:v for k,v in expected.items() if k!='password_revision'}}
        check_expected(db,data,'users','id=?',(uid,),('name','username','role','active'))
        updates=[]; params=[]
        if 'role' in data:
            role=data['role']
            if role not in ('admin','employee'): raise APIError(400,'Некорректные права доступа')
            if uid==user['id'] and role!='admin': raise APIError(400,'Нельзя снять права администратора с самого себя')
            updates.append('role=?'); params.append(role)
            # Права администратора не забирают у человека статус сотрудника:
            # он остаётся в личной таблице, графиках, табеле и отчётах.
            if role=='admin': updates.append('is_staff=1')
            if target['role']=='admin' and role=='employee':
                db.execute('DELETE FROM sessions WHERE user_id=?',(uid,))
        if 'active' in data:
            updates.append('active=?'); params.append(1 if data['active'] is True else 0)
            if data['active'] is not True: db.execute('DELETE FROM sessions WHERE user_id=?',(uid,))
        if 'name' in data:
            name=str(data['name']).strip()[:90]
            if len(name)<2: raise APIError(400,'Введите имя')
            updates.append('name=?'); params.append(name)
        if 'username' in data:
            username=str(data['username']).strip().lower()
            if not USERNAME_PATTERN.fullmatch(username):
                raise APIError(400,'Логин: от 3 до 40 символов, латиница, цифры, точка, дефис или подчёркивание')
            updates.append('username=?'); params.append(username)
            if username!=target['username'].lower():
                db.execute('DELETE FROM sessions WHERE user_id=?',(uid,))
        if 'password' in data:
            if len(str(data['password']))<8: raise APIError(400,'Пароль не менее 8 символов')
            updates.append('password_hash=?'); params.append(hashed_password(str(data['password'])))
            db.execute('DELETE FROM sessions WHERE user_id=?',(uid,))
        if not updates: raise APIError(400,'Нет изменений')
        try: db.execute('UPDATE users SET '+','.join(updates)+' WHERE id=?',params+[uid])
        except sqlite3.IntegrityError: raise APIError(409,'Такой логин уже используется')
        detail=f'Обновлён сотрудник {target["name"]}'
        if 'username' in data and username!=target['username'].lower():
            detail+=f' · логин: @{target["username"]} → @{username}'
        if 'role' in data and data['role']!=target['role']: detail+=f' · права: {"администратор" if data["role"]=="admin" else "сотрудник"}'
        log(db,user['id'],'user_updated',detail); db.commit(); self.send_json({'ok':True})
    def add_task(self,db,user,data):
        title=str(data.get('title','')).strip()[:300]; unit=str(data.get('unit','шт.')).strip()[:30] or 'шт.'
        category=str(data.get('category','Другое')).strip()[:60] or 'Другое'
        try: norm=float(data.get('norm') or 0)
        except (ValueError,TypeError): raise APIError(400,'Некорректная норма')
        if len(title)<4 or not 0<=norm<=10000: raise APIError(400,'Введите название и корректную норму')
        try: res=db.execute('INSERT INTO tasks(title,unit,norm,category) VALUES(?,?,?,?)',(title,unit,norm,category))
        except sqlite3.IntegrityError: raise APIError(409,'Такая работа уже есть в каталоге')
        log(db,user['id'],'task_added',f'Добавлена работа: {title[:120]}'); db.commit(); self.send_json({'ok':True,'id':res.lastrowid})
    def edit_task(self,db,user,path,data):
        try: tid=int(path.rsplit('/',1)[1])
        except ValueError: raise APIError(400,'Некорректная работа')
        old=db.execute('SELECT * FROM tasks WHERE id=?',(tid,)).fetchone()
        if not old: raise APIError(404,'Работа не найдена')
        check_expected(db,data,'tasks','id=?',(tid,),('title','unit','norm','category','active'))
        updates=[]; params=[]
        for key,maxlen in [('title',300),('unit',30),('category',60)]:
            if key in data:
                value=str(data[key]).strip()[:maxlen]
                if not value: raise APIError(400,'Поле не может быть пустым')
                updates.append(key+'=?'); params.append(value)
        if 'norm' in data:
            try: norm=float(data['norm'])
            except (ValueError,TypeError): raise APIError(400,'Некорректная норма')
            if not 0<=norm<=10000: raise APIError(400,'Некорректная норма')
            updates.append('norm=?'); params.append(norm)
        if 'active' in data: updates.append('active=?'); params.append(1 if data['active'] is True else 0)
        if not updates: raise APIError(400,'Нет изменений')
        try: db.execute('UPDATE tasks SET '+','.join(updates)+' WHERE id=?',params+[tid])
        except sqlite3.IntegrityError: raise APIError(409,'Такая работа уже существует')
        log(db,user['id'],'task_updated',f'Изменена работа: {old["title"][:120]}'); db.commit(); self.send_json({'ok':True})
    def password(self,db,user,data):
        if user['role']!='admin': raise APIError(403,'Пароли меняет только администратор')
        old=str(data.get('old','')); new=str(data.get('new',''))
        row=db.execute('SELECT password_hash FROM users WHERE id=?',(user['id'],)).fetchone()
        if not valid_password(row['password_hash'],old): raise APIError(400,'Текущий пароль неверен')
        if len(new)<8: raise APIError(400,'Новый пароль не менее 8 символов')
        db.execute('UPDATE users SET password_hash=? WHERE id=?',(hashed_password(new),user['id']))
        log(db,user['id'],'password','Пароль изменён'); db.commit(); self.send_json({'ok':True})
    def call_cleanup(self,db):
        global CALL_LAST_CLEANUP
        if time.monotonic()-CALL_LAST_CLEANUP<20:return
        # Browser disconnects don't deliver leave reliably; expire heartbeats.
        db.execute("DELETE FROM call_participants WHERE last_seen<datetime('now','-45 seconds')")
        db.execute("UPDATE call_sessions SET state='ended',updated_at=CURRENT_TIMESTAMP WHERE state='active' AND ("
                   "(EXISTS(SELECT 1 FROM call_participants p WHERE p.call_id=call_sessions.id)=0 "
                   "AND created_at<datetime('now','-120 seconds')) OR "
                   "(updated_at<datetime('now','-120 seconds') AND NOT EXISTS(SELECT 1 FROM call_participants p WHERE p.call_id=call_sessions.id)))")
        # A started session ends as soon as its final participant leaves (handled in call_leave).
        db.execute("DELETE FROM call_signals WHERE created_at<datetime('now','-10 minutes')")
        db.execute("DELETE FROM call_sessions WHERE state='ended' AND updated_at<datetime('now','-1 day') AND NOT EXISTS(SELECT 1 FROM call_participants p WHERE p.call_id=call_sessions.id) AND NOT EXISTS(SELECT 1 FROM call_signals s WHERE s.call_id=call_sessions.id)")
        db.commit()
        CALL_LAST_CLEANUP=time.monotonic()
    def call_status(self,db,user,q):
        self.call_cleanup(db)
        calls=[]
        for row in db.execute("SELECT id,room_key,mode,created_by,created_at FROM call_sessions WHERE state='active' ORDER BY id DESC LIMIT 100"):
            room=row['room_key']
            try:
                if room.startswith('dm:'):
                    pair=room[3:].split(':')
                    if len(pair)!=2 or user['id'] not in [int(x) for x in pair]:continue
                else:call_room_key(db,user,room)
            except (ValueError,APIError):continue
            count=db.execute('SELECT COUNT(*) FROM call_participants WHERE call_id=?',(row['id'],)).fetchone()[0]
            calls.append({'id':row['id'],'room':room,'mode':row['mode'],'created_by':row['created_by'],
                          'created_at':row['created_at'],'count':count})
        return self.send_json({'calls':calls,'livekit_available':livekit_ready()})
    def call_start(self,db,user,data):
        room=call_room_key(db,user,data.get('room'))
        self.call_cleanup(db)
        existing=db.execute("SELECT id,mode FROM call_sessions WHERE room_key=? AND state='active' ORDER BY id DESC LIMIT 1",(room,)).fetchone()
        if existing:return self.send_json({'id':existing['id'],'mode':existing['mode'],'room':room,'existing':True})
        mode='livekit' if livekit_ready() else 'p2p'
        sql="INSERT INTO call_sessions(room_key,mode,created_by) VALUES(?,?,?)"
        cursor=db.query(sql,(room,mode,user['id'])) if isinstance(db,d1_store.D1Connection) else db.execute(sql,(room,mode,user['id']))
        db.commit();return self.send_json({'id':cursor.lastrowid,'mode':mode,'room':room,'existing':False})
    def call_cancel(self,db,user,data):
        call=self.call_authorized(db,user,data.get('id'))
        if call['created_by']!=user['id']:raise APIError(403,'Отменить приглашение может только создатель звонка')
        if db.execute('SELECT 1 FROM call_participants WHERE call_id=? LIMIT 1',(call['id'],)).fetchone():
            raise APIError(409,'Другие участники уже подключились')
        db.execute("UPDATE call_sessions SET state='ended',updated_at=CURRENT_TIMESTAMP WHERE id=?",(call['id'],))
        db.commit();return self.send_json({'ok':True})
    def call_authorized(self,db,user,call_id):
        if type(call_id)!=int or call_id<=0:raise APIError(400,'Некорректный звонок')
        row=db.execute('SELECT * FROM call_sessions WHERE id=? AND state=?',(call_id,'active')).fetchone()
        if not row:raise APIError(404,'Звонок завершён')
        room=row['room_key']
        if room.startswith('dm:'):
            pair=room[3:].split(':')
            if len(pair)!=2 or user['id'] not in [int(x) for x in pair]:raise APIError(403,'Нет доступа к звонку')
        else:call_room_key(db,user,room)
        return row
    def call_peer(self,db,user,data):
        call=self.call_authorized(db,user,data.get('id'))
        conn=data.get('connection')
        if not isinstance(conn,str) or not re.fullmatch(r'[0-9a-f]{32}',conn):raise APIError(403,'Сначала подключитесь к звонку')
        if not db.execute('SELECT 1 FROM call_participants WHERE call_id=? AND user_id=? AND connection_id=?',
                          (call['id'],user['id'],conn)).fetchone():raise APIError(403,'Участие в звонке не найдено')
        return call,conn
    def call_join(self,db,user,data):
        self.call_cleanup(db)
        call=self.call_authorized(db,user,data.get('id'))
        count=db.execute('SELECT COUNT(*) FROM call_participants WHERE call_id=?',(call['id'],)).fetchone()[0]
        max_count=15 if call['mode']=='livekit' else 4
        if count>=max_count:raise APIError(409,'В этом режиме звонок заполнен (максимум '+str(max_count)+')')
        conn=secrets.token_hex(16)
        db.execute('INSERT INTO call_participants(call_id,user_id,connection_id) VALUES(?,?,?)',(call['id'],user['id'],conn))
        db.execute('UPDATE call_sessions SET updated_at=CURRENT_TIMESTAMP WHERE id=?',(call['id'],))
        db.commit()
        result={'id':call['id'],'room':call['room_key'],'mode':call['mode'],'connection':conn,
                'max_participants':max_count}
        if call['mode']=='livekit':result.update({'url':os.environ['LIVEKIT_URL'],'token':livekit_token(call,user,conn)})
        return self.send_json(result)
    def call_fallback(self,db,user,data):
        call,conn=self.call_peer(db,user,data)
        if call['mode']!='livekit':raise APIError(400,'Звонок уже работает в резервном режиме')
        count=db.execute('SELECT COUNT(*) FROM call_participants WHERE call_id=?',(call['id'],)).fetchone()[0]
        if count!=1:raise APIError(409,'В LiveKit уже подключились другие участники; переход возможен только до начала встречи')
        db.execute("UPDATE call_sessions SET mode='p2p',updated_at=CURRENT_TIMESTAMP WHERE id=?",(call['id'],))
        db.commit();return self.send_json({'ok':True,'mode':'p2p','max_participants':4})
    def call_heartbeat(self,db,user,data):
        call,conn=self.call_peer(db,user,data)
        db.execute('UPDATE call_participants SET last_seen=CURRENT_TIMESTAMP WHERE call_id=? AND connection_id=?',(call['id'],conn))
        db.execute('UPDATE call_sessions SET updated_at=CURRENT_TIMESTAMP WHERE id=?',(call['id'],))
        db.commit()
        peers=json_rows(db.execute('''SELECT p.connection_id connection,p.user_id,u.name
            FROM call_participants p JOIN users u ON u.id=p.user_id WHERE p.call_id=? AND p.connection_id<>?
            ORDER BY p.connection_id''',(call['id'],conn)).fetchall())
        return self.send_json({'ok':True,'peers':peers})
    def call_leave(self,db,user,data):
        call,conn=self.call_peer(db,user,data)
        db.execute('DELETE FROM call_participants WHERE call_id=? AND connection_id=?',(call['id'],conn))
        db.execute('DELETE FROM call_signals WHERE call_id=? AND (target_conn=? OR sender_conn=?)',(call['id'],conn,conn))
        db.commit()
        if not db.execute('SELECT 1 FROM call_participants WHERE call_id=? LIMIT 1',(call['id'],)).fetchone():
            db.execute("UPDATE call_sessions SET state='ended',updated_at=CURRENT_TIMESTAMP WHERE id=?",(call['id'],))
            db.execute('DELETE FROM call_signals WHERE call_id=?',(call['id'],))
            db.commit()
        return self.send_json({'ok':True})
    def call_signals(self,db,user,q):
        try:call_id=int(q.get('id',[''])[0]);after=int(q.get('after',['0'])[0])
        except (TypeError,ValueError):raise APIError(400,'Некорректная граница сигналов')
        call,conn=self.call_peer(db,user,{'id':call_id,'connection':q.get('connection',[''])[0]})
        if call['mode']!='p2p':raise APIError(400,'Сигналы доступны только для P2P')
        rows=json_rows(db.execute('SELECT id,sender_conn sender,signal_json FROM call_signals WHERE call_id=? AND target_conn=? AND id>? ORDER BY id LIMIT 100',(call_id,conn,max(0,after))).fetchall())
        for row in rows:row['signal']=json.loads(row.pop('signal_json'))
        return self.send_json({'signals':rows})
    def call_signal_send(self,db,user,data):
        call,conn=self.call_peer(db,user,data)
        if call['mode']!='p2p':raise APIError(400,'Недопустимый режим сигнализации')
        target=data.get('target');signal=data.get('signal')
        if not isinstance(target,str) or not re.fullmatch(r'[0-9a-f]{32}',target) or target==conn or not isinstance(signal,dict) or signal.get('type') not in ('offer','answer','candidate'):
            raise APIError(400,'Некорректный сигнал')
        raw=json.dumps(signal,separators=(',',':'))
        if len(raw)>25000 or not db.execute('SELECT 1 FROM call_participants WHERE call_id=? AND connection_id=?',(call['id'],target)).fetchone():
            raise APIError(400,'Участник не найден или сигнал слишком велик')
        db.execute('INSERT INTO call_signals(call_id,sender_conn,target_conn,signal_json) VALUES(?,?,?,?)',(call['id'],conn,target,raw))
        db.commit();return self.send_json({'ok':True})
    def chat_group_create(self,db,user,data):
        title=data.get('title');members=data.get('members')
        if not isinstance(title,str) or not 2<=len(title.strip())<=80 or any(ord(c)<32 for c in title):
            raise APIError(400,'Название беседы: от 2 до 80 символов')
        if not isinstance(members,list) or any(type(x)!=int for x in members) or len(members)!=len(set(members)) or not 3<=len(members)<=49 or any(x==user['id'] or x<=0 for x in members):
            raise APIError(400,'Выберите минимум трёх других сотрудников (всего от четырёх участников)')
        existing=db.execute('SELECT COUNT(*) FROM users WHERE active=1 AND id IN ('+','.join('?' for _ in members)+')',members).fetchone()[0]
        if existing!=len(members):raise APIError(400,'Один из участников не найден или отключён')
        insert='INSERT INTO group_chats(title,created_by) VALUES(?,?)'
        cursor=db.query(insert,(title.strip(),user['id'])) if isinstance(db,d1_store.D1Connection) else db.execute(insert,(title.strip(),user['id']))
        gid=cursor.lastrowid
        for uid in [user['id'],*members]:db.execute('INSERT INTO group_members(group_id,user_id) VALUES(?,?)',(gid,uid))
        log(db,user['id'],'group_created',f'Создана беседа «{title.strip()}» ({len(members)+1} участников)')
        db.commit();return self.send_json({'ok':True,'id':gid,'room':'group:'+str(gid)})
    def chat_group_edit(self,db,user,path,data):
        try:gid=int(path.rsplit('/',1)[1])
        except ValueError:raise APIError(400,'Некорректная беседа')
        room='group:'+str(gid);group=checked_group(db,user,room)
        if group['created_by']!=user['id']:raise APIError(403,'Состав меняет создатель беседы')
        members=data.get('members');title=data.get('title',group['title'])
        if not isinstance(title,str) or not 2<=len(title.strip())<=80 or any(ord(c)<32 for c in title) or not isinstance(members,list) or any(type(x)!=int for x in members) or len(members)!=len(set(members)) or not 3<=len(members)<=49 or any(x<=0 or x==user['id'] for x in members):
            raise APIError(400,'Укажите название и минимум трёх других участников')
        if db.execute('SELECT COUNT(*) FROM users WHERE active=1 AND id IN ('+','.join('?' for _ in members)+')',members).fetchone()[0]!=len(members):
            raise APIError(400,'Участник отключён или не найден')
        old={r[0] for r in db.execute('SELECT user_id FROM group_members WHERE group_id=?',(gid,))}
        deleted=old-set(members)-{user['id']}
        if deleted:
            calls=db.execute("SELECT id FROM call_sessions WHERE room_key=? AND mode='livekit' AND state='active'",(room,)).fetchall()
            for call in calls:
                for member_id in deleted:
                    connections=db.execute('SELECT connection_id FROM call_participants WHERE call_id=? AND user_id=?',
                                           (call['id'],member_id)).fetchall()
                    for connection in connections:livekit_evict(call['id'],connection['connection_id'])
        for uid in deleted:
            # Remove the revoked member's signalling presence immediately.
            # Other honest P2P clients close that peer on the next heartbeat.
            affected=db.execute("SELECT call_id,connection_id FROM call_participants WHERE user_id=? AND call_id IN (SELECT id FROM call_sessions WHERE room_key=? AND state='active')",(uid,room)).fetchall()
            for participant in affected:
                db.execute('DELETE FROM call_signals WHERE call_id=? AND (sender_conn=? OR target_conn=?)',
                           (participant['call_id'],participant['connection_id'],participant['connection_id']))
                db.execute('DELETE FROM call_participants WHERE call_id=? AND connection_id=?',
                           (participant['call_id'],participant['connection_id']))
                db.execute("UPDATE call_sessions SET state='ended',updated_at=CURRENT_TIMESTAMP WHERE id=? AND NOT EXISTS(SELECT 1 FROM call_participants WHERE call_id=?)",
                           (participant['call_id'],participant['call_id']))
            db.execute('DELETE FROM group_members WHERE group_id=? AND user_id=?',(gid,uid))
            db.execute('DELETE FROM chat_pins WHERE user_id=? AND room_key=?',(uid,room))
            db.execute('DELETE FROM chat_reads WHERE user_id=? AND room_key=?',(uid,room))
        for uid in set(members)-old:db.execute('INSERT INTO group_members(group_id,user_id) VALUES(?,?)',(gid,uid))
        db.execute('UPDATE group_chats SET title=? WHERE id=?',(title.strip(),gid))
        log(db,user['id'],'group_updated',f'Обновлён состав беседы #{gid}')
        db.commit();return self.send_json({'ok':True,'room':room})
    def chat_list(self,db,user,q):
        """Only group messages or direct messages involving this user; last 100 + polling."""
        room=q.get('room',['general'])[0]
        if isinstance(room,str) and room.startswith('group:'):
            return self.chat_group_list(db,user,room,q)
        if room=='general':
            where='m.recipient_id IS NULL';params=[]
        else:
            try: peer=int(room)
            except (TypeError,ValueError): raise APIError(400,'Некорректный собеседник')
            if peer<=0 or peer==user['id'] or str(peer)!=room:
                raise APIError(400,'Некорректный собеседник')
            if not db.execute('SELECT id FROM users WHERE id=?',(peer,)).fetchone():
                raise APIError(404,'Собеседник не найден')
            where='((m.sender_id=? AND m.recipient_id=?) OR (m.sender_id=? AND m.recipient_id=?))'
            params=[user['id'],peer,peer,user['id']]
        after=q.get('after',[None])[0];before=q.get('before',[None])[0]
        if after is not None and before is not None: raise APIError(400,'Выберите одну границу истории')
        if after is None:
            older_params=list(params)
            if before is not None:
                try: boundary=int(before)
                except (TypeError,ValueError): raise APIError(400,'Некорректный номер сообщения')
                if boundary<=0 or str(boundary)!=before: raise APIError(400,'Некорректный номер сообщения')
                older_params.append(boundary)
            condition=' AND m.id<?' if before is not None else ''
            sql='''SELECT m.id,m.sender_id,m.recipient_id,m.body,m.created_at,m.style_json,u.name sender_name,
                     COALESCE(a.revision,0) avatar_revision,COALESCE(a.image_url,'') avatar_url,
                     CASE WHEN a.image_b64<>'' THEN 1 ELSE 0 END avatar_stored
                     FROM chat_messages m JOIN users u ON u.id=m.sender_id
                     LEFT JOIN user_avatars a ON a.user_id=u.id
                     WHERE '''+where+condition+' ORDER BY m.id DESC LIMIT 100'
            messages=json_rows(db.execute(sql,older_params).fetchall())[::-1]
            if messages:
                check='SELECT 1 FROM chat_messages m WHERE '+where+' AND m.id<? LIMIT 1'
                has_more=bool(db.execute(check,params+[messages[0]['id']]).fetchone())
            else: has_more=False
        else:
            try: offset=int(after)
            except (TypeError,ValueError): raise APIError(400,'Некорректный номер сообщения')
            if offset<0 or str(offset)!=after: raise APIError(400,'Некорректный номер сообщения')
            sql='''SELECT m.id,m.sender_id,m.recipient_id,m.body,m.created_at,m.style_json,u.name sender_name,
                     COALESCE(a.revision,0) avatar_revision,COALESCE(a.image_url,'') avatar_url,
                     CASE WHEN a.image_b64<>'' THEN 1 ELSE 0 END avatar_stored
                     FROM chat_messages m JOIN users u ON u.id=m.sender_id
                     LEFT JOIN user_avatars a ON a.user_id=u.id
                     WHERE '''+where+' AND m.id>? ORDER BY m.id LIMIT 100'
            messages=json_rows(db.execute(sql,params+[offset]).fetchall())
            has_more=None
        people=json_rows(db.execute('''SELECT u.id,u.name,u.role,COALESCE(a.revision,0) avatar_revision,
                COALESCE(a.image_url,'') avatar_url,
                CASE WHEN a.image_b64<>'' THEN 1 ELSE 0 END avatar_stored
                FROM users u LEFT JOIN user_avatars a ON a.user_id=u.id
                WHERE u.active=1 AND u.id<>? ORDER BY u.name''',(user['id'],)).fetchall())
        for message in messages:
            try: message['style']=valid_chat_style(json.loads(message.pop('style_json') or '{}'))
            except (ValueError,TypeError,APIError):message['style']={}
        key=chat_theme_key(db,user,room)
        theme_row=db.execute('SELECT theme_json,updated_at FROM chat_themes WHERE room_key=?',(key,)).fetchone()
        try: theme=valid_chat_style(json.loads(theme_row['theme_json'])) if theme_row else {}
        except (ValueError,TypeError,APIError):theme={}
        self.send_json({'messages':messages,'users':people,'has_more':has_more,
                        'theme':theme,'theme_updated_at':theme_row['updated_at'] if theme_row else None})
    def chat_group_list(self,db,user,room,q):
        group=checked_group(db,user,room)
        after=q.get('after',[None])[0];before=q.get('before',[None])[0]
        if after is not None and before is not None:raise APIError(400,'Выберите одну границу истории')
        bound=after if after is not None else before
        if bound is not None:
            try:num=int(bound)
            except (TypeError,ValueError):raise APIError(400,'Некорректный номер сообщения')
            if num<0 or str(num)!=bound:raise APIError(400,'Некорректный номер сообщения')
        params=[group['id']]
        if after is not None:
            clause=' AND m.id>?';params.append(num);order='ASC'
        else:
            clause=' AND m.id<?' if before is not None else ''
            if before is not None:params.append(num)
            order='DESC'
        rows=db.execute('''SELECT m.id,m.sender_id,NULL recipient_id,m.body,m.created_at,m.style_json,
               u.name sender_name,COALESCE(a.revision,0) avatar_revision,COALESCE(a.image_url,'') avatar_url,
               CASE WHEN a.image_b64<>'' THEN 1 ELSE 0 END avatar_stored
               FROM group_messages m JOIN users u ON u.id=m.sender_id
               LEFT JOIN user_avatars a ON a.user_id=u.id WHERE m.group_id=?'''+clause+
               ' ORDER BY m.id '+order+' LIMIT 100',params).fetchall()
        messages=json_rows(rows)
        if after is None:messages.reverse()
        has_more=(bool(db.execute('SELECT 1 FROM group_messages WHERE group_id=? AND id<? LIMIT 1',
                       (group['id'],messages[0]['id'])).fetchone()) if messages else False) if after is None else None
        for m in messages:
            try:m['style']=valid_chat_style(json.loads(m.pop('style_json') or '{}'))
            except (ValueError,TypeError,APIError):m['style']={}
        people=json_rows(db.execute('''SELECT u.id,u.name,u.role,COALESCE(a.revision,0) avatar_revision,
             COALESCE(a.image_url,'') avatar_url,CASE WHEN a.image_b64<>'' THEN 1 ELSE 0 END avatar_stored
             FROM users u LEFT JOIN user_avatars a ON a.user_id=u.id WHERE u.active=1 AND u.id<>? ORDER BY u.name''',
             (user['id'],)).fetchall())
        members=[r[0] for r in db.execute('SELECT user_id FROM group_members WHERE group_id=? ORDER BY user_id',(group['id'],))]
        theme_row=db.execute('SELECT theme_json,updated_at FROM chat_themes WHERE room_key=?',(room,)).fetchone()
        try:theme=valid_chat_style(json.loads(theme_row['theme_json'])) if theme_row else {}
        except (ValueError,TypeError,APIError):theme={}
        return self.send_json({'messages':messages,'users':people,'has_more':has_more,'group':{
                  'id':group['id'],'title':group['title'],'created_by':group['created_by'],'members':members},
                  'theme':theme,'theme_updated_at':theme_row['updated_at'] if theme_row else None})
    def chat_theme_save(self,db,user,data):
        key=chat_theme_key(db,user,data.get('room'))
        theme=valid_chat_style(data.get('style'))
        if 'expected' in data:
            current=db.execute('SELECT theme_json FROM chat_themes WHERE room_key=?',(key,)).fetchone()
            stored=valid_chat_style(json.loads(current['theme_json'])) if current else {}
            if stored!=data['expected']:
                raise APIError(409,'Рамку беседы уже изменил другой участник. Обновите чат перед сохранением')
        db.execute('''INSERT INTO chat_themes(room_key,theme_json,updated_by)
            VALUES(?,?,?) ON CONFLICT(room_key) DO UPDATE SET
            theme_json=excluded.theme_json,updated_by=excluded.updated_by,
            updated_at=CURRENT_TIMESTAMP''',
            (key,json.dumps(theme,separators=(',',':')),user['id']))
        db.commit()
        self.send_json({'ok':True,'theme':theme})
    def chat_send(self,db,user,data):
        if data.get('group_id') is not None:
            gid=data.get('group_id')
            if type(gid)!=int or gid<=0:raise APIError(400,'Некорректная беседа')
            checked_group(db,user,'group:'+str(gid))
            text=data.get('body')
            if not isinstance(text,str) or not text.strip() or len(text.strip())>2000 or '\x00' in text:
                raise APIError(400,'Сообщение: от 1 до 2000 символов')
            style=valid_chat_style(data.get('style',{}))
            sql='INSERT INTO group_messages(group_id,sender_id,body,style_json) VALUES(?,?,?,?)'
            args=(gid,user['id'],text.strip(),json.dumps(style,separators=(',',':')))
            cursor=db.query(sql,args) if isinstance(db,d1_store.D1Connection) else db.execute(sql,args)
            db.commit();return self.send_json({'ok':True,'id':cursor.lastrowid})
        recipient=data.get('recipient_id')
        if recipient is not None:
            if type(recipient) is not int or recipient<=0 or recipient==user['id']:
                raise APIError(400,'Некорректный собеседник')
            if not db.execute('SELECT id FROM users WHERE id=? AND active=1',(recipient,)).fetchone():
                raise APIError(404,'Собеседник не найден или отключён')
        text=data.get('body')
        if not isinstance(text,str): raise APIError(400,'Введите сообщение')
        text=text.strip()
        if not text or len(text)>2000 or '\x00' in text:
            raise APIError(400,'Сообщение: от 1 до 2000 символов')
        style=valid_chat_style(data.get('style',{}))
        cursor=db.execute('INSERT INTO chat_messages(sender_id,recipient_id,body,style_json) VALUES(?,?,?,?)',
                          (user['id'],recipient,text,json.dumps(style,separators=(',',':'))))
        db.commit()
        self.send_json({'ok':True,'id':cursor.lastrowid})
    def chat_pin(self,db,user,data):
        room=data.get('room');pin=data.get('pinned')
        if type(pin)!=bool:raise APIError(400,'Некорректная отметка закрепления')
        if isinstance(room,str) and room.startswith('group:'):
            key='group:'+str(checked_group(db,user,room)['id'])
        elif room=='general':key='general'
        else:
            try:peer=int(room)
            except (TypeError,ValueError):raise APIError(400,'Некорректный чат')
            if peer<=0 or peer==user['id'] or str(peer)!=room or not db.execute('SELECT 1 FROM users WHERE id=? AND active=1',(peer,)).fetchone():
                raise APIError(400,'Чат не найден')
            key='dm:'+room
        if 'expected_pinned' in data:
            current=bool(db.execute('SELECT 1 FROM chat_pins WHERE user_id=? AND room_key=?',(user['id'],key)).fetchone())
            if type(data['expected_pinned'])!=bool:raise APIError(400,'Некорректная версия закрепления')
            if current!=data['expected_pinned']:raise APIError(409,'Закрепление чата уже изменено в другой вкладке. Обновите список чатов')
        if pin:db.execute('INSERT OR IGNORE INTO chat_pins(user_id,room_key) VALUES(?,?)',(user['id'],key))
        else:db.execute('DELETE FROM chat_pins WHERE user_id=? AND room_key=?',(user['id'],key))
        db.commit();self.send_json({'ok':True,'room':room,'pinned':pin})
    def chat_status(self,db,user,q):
        """Unread counters and new-message notices; first poll never replays old toasts."""
        uid=user['id']
        direct=json_rows(db.execute('''SELECT m.sender_id peer_id,COUNT(*) unread
            FROM chat_messages m LEFT JOIN chat_reads r
              ON r.user_id=? AND r.room_key=('dm:' || m.sender_id)
            WHERE m.recipient_id=? AND m.id>COALESCE(r.last_read_id,0)
            GROUP BY m.sender_id''',(uid,uid)).fetchall())
        general=db.execute('''SELECT COUNT(*) FROM chat_messages m LEFT JOIN chat_reads r
              ON r.user_id=? AND r.room_key='general'
            WHERE m.recipient_id IS NULL AND m.sender_id<>?
              AND m.id>COALESCE(r.last_read_id,0)''',(uid,uid)).fetchone()[0]
        latest=db.execute('SELECT COALESCE(MAX(id),0) FROM chat_messages').fetchone()[0]
        after=q.get('after',[None])[0]
        notices=[]
        if after is not None:
            try: cursor=int(after)
            except (TypeError,ValueError): raise APIError(400,'Некорректный номер сообщения')
            if cursor<0 or cursor>9223372036854775807 or str(cursor)!=after:
                raise APIError(400,'Некорректный номер сообщения')
            notices=json_rows(db.execute('''SELECT m.id,m.sender_id,m.recipient_id,u.name sender_name
                FROM chat_messages m JOIN users u ON u.id=m.sender_id
                WHERE m.id>? AND m.sender_id<>?
                  AND (m.recipient_id=? OR m.recipient_id IS NULL)
                ORDER BY m.id LIMIT 50''',(cursor,uid,uid)).fetchall())
        rooms=db.execute("""SELECT CASE WHEN m.recipient_id IS NULL THEN 'general'
                   WHEN m.sender_id=? THEN 'dm:'||m.recipient_id ELSE 'dm:'||m.sender_id END room_key,
                   MAX(m.id) last_id FROM chat_messages m
                   WHERE m.recipient_id IS NULL OR m.sender_id=? OR m.recipient_id=?
                   GROUP BY room_key""",(uid,uid,uid)).fetchall()
        pinned=db.execute('SELECT room_key FROM chat_pins WHERE user_id=?',(uid,)).fetchall()
        groups=json_rows(db.execute('''SELECT g.id,g.title,g.created_by,
             (SELECT COUNT(*) FROM group_members members WHERE members.group_id=g.id) member_count,
             (SELECT COALESCE(MAX(id),0) FROM group_messages m WHERE m.group_id=g.id) last_id,
             (SELECT COUNT(*) FROM group_messages m LEFT JOIN chat_reads r
                ON r.user_id=? AND r.room_key=('group:'||g.id)
                WHERE m.group_id=g.id AND m.sender_id<>? AND m.id>COALESCE(r.last_read_id,0)) unread
             FROM group_chats g JOIN group_members gm ON gm.group_id=g.id
             WHERE gm.user_id=? ORDER BY g.id DESC''',(uid,uid,uid)).fetchall())
        group_latest=db.execute('''SELECT COALESCE(MAX(m.id),0) FROM group_messages m
            JOIN group_members gm ON gm.group_id=m.group_id AND gm.user_id=?''',(uid,)).fetchone()[0]
        group_after=q.get('group_after',[None])[0]
        group_notices=[]
        if group_after is not None:
            try:group_cursor=int(group_after)
            except (TypeError,ValueError):raise APIError(400,'Некорректный номер группового сообщения')
            if group_cursor<0 or str(group_cursor)!=group_after:raise APIError(400,'Некорректный номер группового сообщения')
            group_notices=json_rows(db.execute('''SELECT m.id,m.group_id,m.sender_id,u.name sender_name
                FROM group_messages m JOIN group_members gm ON gm.group_id=m.group_id AND gm.user_id=?
                JOIN users u ON u.id=m.sender_id WHERE m.id>? AND m.sender_id<>?
                ORDER BY m.id LIMIT 50''',(uid,group_cursor,uid)).fetchall())
        next_group_cursor=group_notices[-1]['id'] if len(group_notices)==50 else group_latest
        # If there are more than 50 notices, the next poll continues instead of skipping them.
        next_cursor=notices[-1]['id'] if len(notices)==50 else latest
        self.send_json({'direct':{str(x['peer_id']):x['unread'] for x in direct},
                        'general':general,'notices':notices,'cursor':next_cursor,
                        'rooms':{r['room_key']:r['last_id'] for r in rooms},
                        'pinned':[r['room_key'] for r in pinned],
                        'groups':groups,'group_notices':group_notices,'group_cursor':next_group_cursor})
    def chat_read(self,db,user,data):
        room=data.get('room')
        if isinstance(room,str) and room.startswith('group:'):
            group=checked_group(db,user,room);key='group:'+str(group['id'])
            latest=db.execute('SELECT COALESCE(MAX(id),0) FROM group_messages WHERE group_id=?',(group['id'],)).fetchone()[0]
        elif room=='general':
            key='general';latest=db.execute('''SELECT COALESCE(MAX(id),0) FROM chat_messages
                WHERE recipient_id IS NULL''').fetchone()[0]
        else:
            try: peer=int(room)
            except (TypeError,ValueError): raise APIError(400,'Некорректный собеседник')
            if peer<=0 or peer==user['id'] or str(peer)!=room or not db.execute('SELECT id FROM users WHERE id=?',(peer,)).fetchone():
                raise APIError(400,'Некорректный собеседник')
            key='dm:'+str(peer)
            latest=db.execute('''SELECT COALESCE(MAX(id),0) FROM chat_messages
                WHERE (sender_id=? AND recipient_id=?) OR (sender_id=? AND recipient_id=?)''',
                (user['id'],peer,peer,user['id'])).fetchone()[0]
        if 'up_to' in data:
            visible=data['up_to']
            if type(visible) is not int or visible<0 or visible>latest:
                raise APIError(400,'Некорректный номер показанного сообщения')
            latest=visible
        db.execute('''INSERT INTO chat_reads(user_id,room_key,last_read_id) VALUES(?,?,?)
            ON CONFLICT(user_id,room_key) DO UPDATE SET
            last_read_id=MAX(chat_reads.last_read_id,excluded.last_read_id)''',(user['id'],key,latest))
        db.commit();self.send_json({'ok':True,'last_read_id':latest})
    def avatar_save(self,db,user,data):
        image='';mime='';url=''
        if data.get('remove') is True:
            pass
        elif 'url' in data:
            url=data['url']
            if not isinstance(url,str) or not 1<=len(url)<=800 or url!=url.strip() or any(ord(c)<32 for c in url):
                raise APIError(400,'Введите корректную HTTPS-ссылку на изображение')
            try:
                parsed=urlparse(url)
                allowed=parsed.scheme=='https' and bool(parsed.hostname) and not parsed.username and not parsed.password and parsed.port in (None,443)
            except ValueError: allowed=False
            if not allowed:
                raise APIError(400,'Разрешены только прямые HTTPS-ссылки без логина и пароля')
            host=parsed.hostname.lower().rstrip('.')
            if '.' not in host or host.endswith(('.local','.internal','.localhost')):
                raise APIError(400,'Укажите публичную HTTPS-ссылку на изображение')
            try:
                address=ipaddress.ip_address(host)
                if not address.is_global: raise APIError(400,'Недоступный адрес изображения')
            except ValueError: pass
        elif 'image_b64' in data:
            encoded=data['image_b64']
            if not isinstance(encoded,str) or not 1<=len(encoded)<=130000:
                raise APIError(413,'Аватар слишком большой (после сжатия — до 90 КБ)')
            try: raw=base64.b64decode(encoded,validate=True)
            except (ValueError,binascii.Error): raise APIError(400,'Некорректный файл изображения')
            if not raw or len(raw)>90000: raise APIError(413,'Аватар слишком большой (до 90 КБ)')
            if raw.startswith(b'\x89PNG\r\n\x1a\n') and raw[12:16]==b'IHDR':mime='image/png'
            elif raw.startswith(b'\xff\xd8\xff'):mime='image/jpeg'
            elif raw[:4]==b'RIFF' and raw[8:12]==b'WEBP':mime='image/webp'
            else:raise APIError(400,'Аватар должен быть PNG, JPEG или WebP')
            image=base64.b64encode(raw).decode('ascii')
        else:raise APIError(400,'Выберите файл изображения или укажите HTTPS-ссылку')
        if 'expected_revision' in data:
            previous=db.execute('SELECT revision FROM user_avatars WHERE user_id=?',(user['id'],)).fetchone()
            if type(data['expected_revision'])!=int or data['expected_revision']!=(previous[0] if previous else 0):
                raise APIError(409,'Фото профиля уже изменилось в другой вкладке. Обновите страницу')
        db.execute('''INSERT INTO user_avatars(user_id,image_b64,mime,image_url,revision)
            VALUES(?,?,?,?,1) ON CONFLICT(user_id) DO UPDATE SET
            image_b64=excluded.image_b64,mime=excluded.mime,image_url=excluded.image_url,
            revision=user_avatars.revision+1''',(user['id'],image,mime,url))
        db.commit()
        revision=db.execute('SELECT revision FROM user_avatars WHERE user_id=?',(user['id'],)).fetchone()[0]
        self.send_json({'ok':True,'avatar_revision':revision,'avatar_url':url,'avatar_stored':bool(image)})
    def avatar_image(self,db,user,path):
        try: uid=int(path.rsplit('/',1)[1])
        except (TypeError,ValueError): raise APIError(400,'Некорректный пользователь')
        if uid<=0 or str(uid)!=path.rsplit('/',1)[1]: raise APIError(400,'Некорректный пользователь')
        row=db.execute('SELECT image_b64,mime FROM user_avatars WHERE user_id=?',(uid,)).fetchone()
        if not row or not row['image_b64']: raise APIError(404,'Аватар не найден')
        raw=base64.b64decode(row['image_b64'])
        self.send_response(200)
        self.send_header('Content-Type',row['mime'])
        self.send_header('Content-Length',str(len(raw)))
        self.send_header('Cache-Control','private, no-store')
        self.send_header('X-Content-Type-Options','nosniff')
        self.end_headers();self.wfile.write(raw)
    def report_data(self,db,month):
        first,last=month_range(month)
        rows=json_rows(db.execute('''SELECT t.id task_id,t.title,t.unit,t.norm,t.category,u.id user_id,u.name,
            ROUND(SUM(e.quantity),3) quantity,ROUND(SUM(e.quantity)*t.norm,3) hours
            FROM entries e JOIN tasks t ON t.id=e.task_id JOIN users u ON u.id=e.user_id
            WHERE e.work_date BETWEEN ? AND ? AND u.role='employee' AND u.is_staff=1
            GROUP BY t.id,u.id ORDER BY t.id,u.name''',(first,last)).fetchall())
        users=json_rows(db.execute("SELECT id,name FROM users WHERE is_staff=1 AND role='employee' ORDER BY name").fetchall())
        tasks=json_rows(db.execute('SELECT id,title,unit,norm,category,active FROM tasks ORDER BY id').fetchall())
        attendance=json_rows(db.execute('''SELECT a.user_id,a.day,a.status,a.reason FROM attendance a
            JOIN users u ON u.id=a.user_id WHERE a.day BETWEEN ? AND ?
            AND u.role='employee' AND u.is_staff=1''',(first,last)).fetchall())
        daily_hours=json_rows(db.execute('SELECT day,hours FROM daily_hours WHERE day BETWEEN ? AND ?',(first,last)).fetchall())
        personal_hours=json_rows(db.execute('''SELECT p.user_id,p.day,p.hours FROM personal_hours p
            JOIN users u ON u.id=p.user_id WHERE p.day BETWEEN ? AND ?
            AND u.role='employee' AND u.is_staff=1''',(first,last)).fetchall())
        comments=json_rows(db.execute('''SELECT c.user_id,c.task_id,c.work_date,c.body,u.name,t.title
            FROM cell_comments c JOIN users u ON u.id=c.user_id JOIN tasks t ON t.id=c.task_id
            WHERE c.work_date BETWEEN ? AND ? AND u.role='employee' AND u.is_staff=1
            ORDER BY c.work_date,u.name''',(first,last)).fetchall())
        return rows,users,tasks,attendance,daily_hours,personal_hours,comments,first,last
    def report(self,db,user,q):
        month=q.get('month',[TODAY().strftime('%Y-%m')])[0]
        rows,users,tasks,attendance,daily_hours,personal_hours,comments,first,last=self.report_data(db,month)
        productivity,overall=calculate_productivity(month,rows,users,attendance,daily_hours,personal_hours)
        self.send_json({'month':month,'rows':rows,'users':users,'tasks':tasks,'attendance':attendance,
            'daily_hours':daily_hours,'personal_hours':personal_hours,'comments':comments,
            'productivity':productivity,'productivity_total':overall})
    def export(self,db,user,q):
        month=q.get('month',[TODAY().strftime('%Y-%m')])[0]
        rows,users,tasks,attendance,daily_hours,personal_hours,comments,first,last=self.report_data(db,month)
        # ID сохраняет порядок сотрудников исходного файла (листы импортировались по порядку).
        # Новые сотрудники размещаются после них; сайт может сортировать их по имени.
        users.sort(key=lambda person:person['id'])
        productivity,overall=calculate_productivity(month,rows,users,attendance,daily_hours,personal_hours)
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font,PatternFill,Alignment,Border,Side
            from openpyxl.utils import get_column_letter
        except ImportError: raise APIError(500,'Для экспорта Excel установите openpyxl')
        wb=Workbook(); ws=wb.active; ws.title='Отчёт'
        # Повторяем сетку вкладки «Отчёт» исходного source.xlsx, не копируя
        # в публичный код сам файл с личными данными сотрудников.
        # A:B — работа/единица; затем количество по людям + общее;
        # норма; затем время по ТЕМ ЖЕ людям + итого.
        count=len(users)
        qty_start=3
        qty_total=qty_start+count
        norm_col=qty_total+1
        time_start=norm_col+1
        time_total=time_start+count
        year=int(month[:4]); selected_month=int(month[5:])
        months=['Январь','Февраль','Март','Апрель','Май','Июнь',
                'Июль','Август','Сентябрь','Октябрь','Ноябрь','Декабрь']
        navy='173347'; ink='222222'; border_color='C9C9C9'; total_fill='F4F1E6'
        thin=Side(style='hair',color=border_color)
        grid=Border(left=thin,right=thin,top=thin,bottom=thin)
        ws.sheet_view.showGridLines=True
        ws.merge_cells(start_row=1,start_column=1,end_row=1,end_column=norm_col)
        c=ws.cell(1,1,'Отчет по выполненной работе в программных комплексах 7ЭГ, МПК "Панорама", ПК "УРГ" за')
        c.font=Font(name='Calibri',size=17,color=ink)
        c.alignment=Alignment(horizontal='right',vertical='center',wrap_text=True)
        if count>1: ws.merge_cells(start_row=1,start_column=time_start,end_row=1,end_column=time_start+1)
        ws.cell(1,time_start,months[selected_month-1]).font=Font(name='Calibri',size=13,bold=True,color=ink)
        ws.cell(1,time_start+2,year).font=Font(name='Calibri',size=13,color=ink)
        ws.cell(1,time_start+3,'г.').font=Font(name='Calibri',size=11,color=ink)
        ws.row_dimensions[1].height=20.9
        ws.merge_cells(start_row=2,start_column=1,end_row=3,end_column=1)
        ws.merge_cells(start_row=2,start_column=2,end_row=3,end_column=2)
        ws.merge_cells(start_row=2,start_column=norm_col,end_row=3,end_column=norm_col)
        ws.merge_cells(start_row=2,start_column=time_total,end_row=3,end_column=time_total)
        if count>1:
            ws.merge_cells(start_row=2,start_column=qty_start,end_row=2,end_column=qty_total-1)
            ws.merge_cells(start_row=2,start_column=time_start,end_row=2,end_column=time_total-1)
        for rownum in (2,3):
            for col in range(1,time_total+1):
                cell=ws.cell(rownum,col)
                cell.border=grid
                cell.font=Font(name='Calibri',size=11,color=ink)
                cell.alignment=Alignment(horizontal='center',vertical='center',wrap_text=True)
        ws.cell(2,1,'Выполняемые работы')
        ws.cell(2,2,'Кол-во выполненной работы')
        if count: ws.cell(2,qty_start,'По количеству')
        ws.cell(3,qty_total,'Общее')
        ws.cell(2,norm_col,'Норма времени')
        if count: ws.cell(2,time_start,'По времени')
        ws.cell(2,time_total,'Итого')
        for i,person in enumerate(users):
            ws.cell(3,qty_start+i,person['name'])
            ws.cell(3,time_start+i,person['name'])
        ws.row_dimensions[2].height=13.8
        ws.row_dimensions[3].height=33.5
        by_task_user={(x['task_id'],x['user_id']):x for x in rows}
        recorded_tasks={x['task_id'] for x in rows}
        displayed_tasks=[x for x in tasks if x['active'] or x['id'] in recorded_tasks]
        num_format='#,##0.###'; hours_format='#,##0.##'
        def excel_label(value):
            # openpyxl считает строку, начинающуюся с «=», формулой.
            text=str(value or '')
            return "'"+text if text.startswith('=') else text
        for i,task in enumerate(displayed_tasks):
            r=4+i
            ws.cell(r,1,excel_label(task['title']))
            ws.cell(r,2,excel_label(task['unit']))
            ws.cell(r,norm_col,float(task['norm']))
            quantity=0.0;work_time=0.0
            for j,person in enumerate(users):
                record=by_task_user.get((task['id'],person['id']))
                if record:
                    q=float(record['quantity']);h=float(record['hours'])
                    ws.cell(r,qty_start+j,q)
                    ws.cell(r,time_start+j,h)
                    quantity+=q;work_time+=h
            ws.cell(r,qty_total,round(quantity,3))
            ws.cell(r,time_total,round(work_time,3))
            for col in range(1,time_total+1):
                c=ws.cell(r,col)
                c.border=grid
                c.font=Font(name='Calibri',size=11 if col in (1,2,norm_col) else 10,color=ink)
                c.alignment=Alignment(horizontal='left' if col==1 else 'center',vertical='center',wrap_text=col==1)
                if col in (qty_total,time_total): c.fill=PatternFill('solid',fgColor=total_fill)
                if col>=qty_start and isinstance(c.value,(int,float)):
                    c.number_format=hours_format if col>=time_start else num_format
            # Высоты тех же 23 исходных строк: длинная первая и последняя,
            # остальные по 23.75 пт, как на вкладке «Отчёт» source.xlsx.
            if i==0: ws.row_dimensions[r].height=95.25
            elif i==22: ws.row_dimensions[r].height=42.75
            else: ws.row_dimensions[r].height=max(23.75,16*(1+len(task['title'])//65))
        total_row=4+len(displayed_tasks)
        percent_row=total_row+1
        by_person={p['user_id']:p for p in productivity}
        total_qty=round(sum(float(x['quantity']) for x in rows),3)
        # В оригинале итог по нормо-часам округляется один раз после суммирования.
        rounded_total=math.floor(float(overall['norm_hours'])+0.5)
        for r in (total_row,percent_row):
            for col in range(1,time_total+1):
                c=ws.cell(r,col)
                c.border=grid
                c.font=Font(name='Calibri',size=11,color=ink)
                c.alignment=Alignment(horizontal='left' if col==1 else 'center',vertical='center')
                if col in (qty_total,time_total): c.fill=PatternFill('solid',fgColor=total_fill)
        ws.cell(total_row,1,'Общее'); ws.cell(total_row,2,'шт.')
        ws.cell(percent_row,1,'Производительность'); ws.cell(percent_row,2,'%')
        ws.cell(total_row,qty_total,total_qty)
        ws.cell(total_row,time_total,rounded_total)
        for i,person in enumerate(users):
            uid=person['id'];p=by_person[uid]
            ws.cell(total_row,qty_start+i,round(sum(float(x['quantity']) for x in rows if x['user_id']==uid),3))
            ws.cell(total_row,time_start+i,p['rounded_norm_hours'])
            ws.cell(percent_row,time_start+i,p['percent']/100 if p['percent'] is not None else '—')
        ws.cell(percent_row,time_total,rounded_total/overall['work_hours'] if overall['work_hours'] else '—')
        for col in range(qty_start,time_total+1):
            ws.cell(total_row,col).number_format=num_format if col<=qty_total else '#,##0'
        for col in range(time_start,time_total+1):ws.cell(percent_row,col).number_format='0.00%'
        ws.row_dimensions[total_row].height=14.25
        ws.row_dimensions[percent_row].height=14.25
        # Под таблицей, как в исходнике, — «Отработано … часов» по каждому человеку.
        ws.row_dimensions[percent_row+1].height=14.25
        hours_start=percent_row+2
        for i,person in enumerate(users):
            r=hours_start+i
            for col in range(1,6):
                c=ws.cell(r,col);c.border=grid;c.alignment=Alignment(horizontal='center',vertical='center')
                c.font=Font(name='Calibri',size=11,color=ink)
            ws.cell(r,1,excel_label(person['name']))
            ws.cell(r,2,'Отработано')
            ws.cell(r,4,by_person[person['id']]['work_hours']).number_format=hours_format
            ws.cell(r,5,'часов')
            ws.row_dimensions[r].height=13.8 if i<8 else 14.25
        # В исходном файле внизу расположена история табельных часов: Май–Декабрь
        # и блок «Всё». Берём факты из базы, а не устаревшие числа source.xlsx.
        history_row=max(50,hours_start+count+10)
        annual_first=f'{year}-05-01'; annual_last=f'{year}-12-31'
        annual_attendance=json_rows(db.execute('SELECT user_id,day,status FROM attendance WHERE day BETWEEN ? AND ?',
                                              (annual_first,annual_last)).fetchall())
        annual_daily=json_rows(db.execute('SELECT day,hours FROM daily_hours WHERE day BETWEEN ? AND ?',
                                         (annual_first,annual_last)).fetchall())
        annual_personal=json_rows(db.execute('SELECT user_id,day,hours FROM personal_hours WHERE day BETWEEN ? AND ?',
                                            (annual_first,annual_last)).fetchall())
        starts=[1,7,16,24,30,36,42,48,54]  # A/G/P/X/AD/AJ/AP/AV/BB в source.xlsx
        hours_offsets=[3,5,3,3,3,3,3,3,3]
        unit_offsets=[4,7,4,4,4,4,4,4,4]
        year_hours={u['id']:0.0 for u in users}
        through=min(selected_month, int(TODAY().strftime('%m')) if year==TODAY().year else (12 if year<TODAY().year else 0))
        for j,m in enumerate(range(5,13)):
            start=starts[j]; hour_col=start+hours_offsets[j]
            ws.cell(history_row,start+1,months[m-1]).font=Font(name='Calibri',size=11,bold=True,color=ink)
            ws.row_dimensions[history_row].height=12.8
            if m<=through:
                key=f'{year}-{m:02d}'
                result,_=calculate_productivity(key,[],users,
                    [x for x in annual_attendance if x['day'].startswith(key)],
                    [x for x in annual_daily if x['day'].startswith(key)],
                    [x for x in annual_personal if x['day'].startswith(key)])
                month_hours={p['user_id']:p['work_hours'] for p in result}
            else: month_hours={}
            for i,person in enumerate(users):
                r=history_row+1+i;uid=person['id']
                ws.cell(r,start,excel_label(person['name']))
                ws.cell(r,start+1,'Отработано')
                if uid in month_hours:
                    ws.cell(r,hour_col,month_hours[uid]).number_format=hours_format
                    year_hours[uid]+=month_hours[uid]
                ws.cell(r,start+unit_offsets[j],'часов')
        all_start=starts[-1];all_hours_col=all_start+hours_offsets[-1]
        ws.cell(history_row,all_start+1,'Всё').font=Font(name='Calibri',size=11,bold=True,color=ink)
        for i,person in enumerate(users):
            r=history_row+1+i;uid=person['id']
            ws.cell(r,all_start,excel_label(person['name']))
            ws.cell(r,all_start+1,'Отработано')
            ws.cell(r,all_hours_col,round(year_hours[uid],2)).number_format=hours_format
            ws.cell(r,all_start+unit_offsets[-1],'часов')
        # Ширины A:Y исходного листа (при 10 сотрудниках); при изменении
        # состава сохраняем ту же пропорцию: работа шире, числовые столбцы узкие.
        source_widths=[34.62,12.67,13.96,12.67,13,13,13.07,12.67,13,13,13,13,
                       13,13,13.52,12.67,13,13,13,13,13,13,13,13,13]
        ws.column_dimensions['A'].width=source_widths[0]
        ws.column_dimensions['B'].width=source_widths[1]
        for col in range(qty_start,time_total+1):
            ws.column_dimensions[get_column_letter(col)].width=source_widths[col-1] if count==10 else 13
        ws.freeze_panes='C4'
        ws.page_setup.orientation='landscape'
        ws.page_setup.paperSize=ws.PAPERSIZE_A3
        ws.sheet_properties.pageSetUpPr.fitToPage=True
        ws.page_setup.fitToWidth=1;ws.page_setup.fitToHeight=0
        ws.print_title_rows='1:3'
        ws.print_area=f'A1:{get_column_letter(time_total)}{hours_start+count-1 if count else percent_row}'
        at=wb.create_sheet('Табель 5-2'); day_count=int(last[-2:])
        at.append(['Сотрудник']+[str(i) for i in range(1,day_count+1)]+['Рабочих дней','Отсутствий','Часов по табелю'])
        for cell in at[1]:cell.fill=PatternFill('solid',fgColor=navy);cell.font=Font(bold=True,color='FFFFFF')
        attmap={(x['user_id'],x['day']):x for x in attendance}
        global_map={x['day']:x['hours'] for x in daily_hours}
        personal_map={(x['user_id'],x['day']):x['hours'] for x in personal_hours}
        for u in users:
            vals=[u['name']]; presents=0; absents=0; total_hours=0
            for i in range(1,day_count+1):
                d=f'{month}-{i:02d}'; weekday=dt.date.fromisoformat(d).weekday()<5
                override=attmap.get((u['id'],d)); status=override['status'] if override else ('present' if weekday else 'off')
                if status=='present':
                    presents+=1
                    hours=personal_map.get((u['id'],d),global_map.get(d,scheduled_hours(d)))
                    total_hours+=hours; vals.append(hours)
                elif status=='absent': absents+=1; vals.append(0)
                else: vals.append('—')
            at.append(vals+[presents,absents,round(total_hours,2)]); rr=at.max_row
            for i in range(1,day_count+1):
                d=f'{month}-{i:02d}'; override=attmap.get((u['id'],d))
                status=override['status'] if override else ('present' if dt.date.fromisoformat(d).weekday()<5 else 'off')
                cell=at.cell(rr,i+1);cell.alignment=Alignment(horizontal='center');cell.number_format='0.##'
                cell.fill=PatternFill('solid',fgColor='DDF3E8' if status=='present' else 'FDE3DF' if status=='absent' else 'F2F4F7')
                if status=='absent':cell.comment=__import__('openpyxl').comments.Comment(override['reason'] or 'Отсутствует','Табель')
            at.cell(rr,at.max_column).number_format='0.##'
        at.column_dimensions['A'].width=26
        for i in range(2,at.max_column+1):at.column_dimensions[get_column_letter(i)].width=9 if i<=day_count+1 else 19
        at.freeze_panes='B2'
        detail=wb.create_sheet('Записи');detail.append(['Дата','Сотрудник','Работа','Количество','Ед. изм.','Примечание'])
        for cell in detail[1]:cell.fill=PatternFill('solid',fgColor=navy);cell.font=Font(bold=True,color='FFFFFF')
        entries=db.execute('''SELECT e.work_date,u.name,t.title,e.quantity,t.unit,e.note FROM entries e JOIN users u ON u.id=e.user_id JOIN tasks t ON t.id=e.task_id WHERE e.work_date BETWEEN ? AND ? AND u.role='employee' AND u.is_staff=1 ORDER BY e.work_date,u.name''',(first,last)).fetchall()
        for entry in entries: detail.append(list(entry))
        for col,width in {'A':17,'B':24,'C':76,'D':18,'E':15,'F':45}.items():detail.column_dimensions[col].width=width
        detail.freeze_panes='C2';detail.auto_filter.ref=f'A1:F{detail.max_row}'
        notes=wb.create_sheet('Комментарии')
        notes.append(['Дата','Сотрудник','Работа','Комментарий'])
        for cell in notes[1]: cell.fill=PatternFill('solid',fgColor=navy);cell.font=Font(bold=True,color='FFFFFF')
        for item in comments: notes.append([item['work_date'],item['name'],item['title'],item['body']])
        for col,width in {'A':17,'B':24,'C':76,'D':80}.items(): notes.column_dimensions[col].width=width
        notes.freeze_panes='C2';notes.auto_filter.ref=f'A1:D{notes.max_row}'
        out=io.BytesIO();wb.save(out)
        self.send_bytes(out.getvalue(),f'forma-report-{month}.xlsx','application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')

if __name__=='__main__':
    init_db()
    if os.environ.get('RENDER')=='true':
        with connect() as db:
            if not db.execute('SELECT id FROM users WHERE role="admin" AND active=1 LIMIT 1').fetchone():
                if len(os.environ.get('INITIAL_IMPORT_TOKEN',''))<32:
                    raise RuntimeError('D1 is empty: set a strong INITIAL_IMPORT_TOKEN for browser setup, or migrate a verified backup first')
                print('D1 is empty. One-time ZIP upload is available at /setup; remove INITIAL_IMPORT_TOKEN after importing.',flush=True)
    port=int(os.environ.get('PORT','8000'))
    print(f'Forma running at http://0.0.0.0:{port}',flush=True)
    ThreadingHTTPServer(('0.0.0.0',port),Handler).serve_forever()
