"""Private vacation schedule data, validation, XLSX import/export; no source sheet in public code."""
import calendar
import datetime as dt
import io
import re
import unicodedata

SCHEMA='''CREATE TABLE IF NOT EXISTS vacation_schedule (
 user_id INTEGER NOT NULL REFERENCES users(id),
 year INTEGER NOT NULL CHECK(year BETWEEN 2020 AND 2100),
 month INTEGER NOT NULL CHECK(month BETWEEN 1 AND 12),
 start_date TEXT NOT NULL, end_date TEXT NOT NULL,
 material_help INTEGER NOT NULL DEFAULT 0 CHECK(material_help IN (0,1)),
 updated_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,
 PRIMARY KEY(user_id,year,month));'''
MONTHS=['Январь','Февраль','Март','Апрель','Май','Июнь','Июль','Август','Сентябрь','Октябрь','Ноябрь','Декабрь']
DATE_RE=re.compile(r'\b(\d{2})\.(\d{2})\.(\d{4})\b')

class VacationError(ValueError):pass

def valid_year(value):
    try: year=int(value)
    except (TypeError,ValueError): raise VacationError('Некорректный год')
    if str(value)!=str(year) or not 2020<=year<=dt.date.today().year+10:
        raise VacationError('Год вне допустимого диапазона')
    return year

def dates(year,month,start,end):
    try:
        if type(month)!=int or not 1<=month<=12: raise ValueError()
        if not isinstance(start,str) or not isinstance(end,str): raise ValueError()
        if not re.fullmatch(r'\d{4}-\d{2}-\d{2}',start) or not re.fullmatch(r'\d{4}-\d{2}-\d{2}',end): raise ValueError()
        first=dt.date.fromisoformat(start);last=dt.date.fromisoformat(end)
        if first.year!=year or first.month!=month or not 0<=(last-first).days<=93: raise ValueError()
    except ValueError: raise VacationError('Укажите корректный период: начало в выбранном месяце, конец не ранее начала (не более 94 дней)')
    return (last-first).days+1

def roster(db,year):
    people=[dict(r) for r in db.execute('SELECT id,name FROM users WHERE is_staff=1 ORDER BY name,id')]
    items=[dict(r) for r in db.execute('''SELECT v.user_id,v.month,v.start_date,v.end_date,v.material_help
      FROM vacation_schedule v JOIN users u ON u.id=v.user_id WHERE v.year=? AND u.is_staff=1
      ORDER BY v.user_id,v.month''',(year,))]
    for row in items:row['days']=dates(year,row['month'],row['start_date'],row['end_date'])
    return people,items

def _key(name):
    return ' '.join(str(name).casefold().replace('ё','е').split())

def parse_xlsx(blob,people,year,overrides=None):
    """Read only top sheet's dates. Lower sheet's +М applies only to an existing upper cell."""
    from openpyxl import load_workbook
    if not blob.startswith(b'PK') or len(blob)>2*1024*1024:raise VacationError('Выберите XLSX до 2 МБ')
    try: book=load_workbook(io.BytesIO(blob),read_only=True,data_only=False)
    except Exception as exc:raise VacationError('Не удалось прочитать XLSX') from exc
    try:
        sheet=book.worksheets[0]
        if (sheet.max_row or 0)>500 or (sheet.max_column or 0)>40:raise VacationError('Файл слишком большой')
        by_name={}
        for person in people:
            by_name.setdefault(_key(person['name']),[]).append(person['id'])
            by_name.setdefault(_key(person['name']).split()[0],[]).append(person['id'])
        found=[];missing=[];entries=[];warnings=[];mapping=[];overrides=overrides or {}
        by_id={p['id']:p for p in people}
        for row in range(3,min(sheet.max_row or 200,200)+1):
            name=sheet.cell(row,1).value
            if not isinstance(name,str) or not name.strip():continue
            # The second part duplicates the same people and is not a new roster.
            if _key(name)=='фио':break
            full=_key(name);surname=full.split()[0]
            matches=list(dict.fromkeys(by_name.get(full,[]) or by_name.get(surname,[])))
            manual=overrides.get(name.strip())
            if manual is not None:
                if type(manual)!=int or (manual!=0 and manual not in by_id):raise VacationError('Сотрудник для строки «'+name.strip()+'» не найден')
                matches=[manual] if manual else []
            uid=matches[0] if len(matches)==1 else None
            mapping.append({'source':name.strip(),'user_id':uid,
                            'target':by_id[uid]['name'] if uid else '',
                            'reason':'Пропущено вручную' if manual==0 else 'Выбрано вручную' if manual is not None else 'Совпадение' if uid else 'Не найден или несколько однофамильцев'})
            if uid is None:
                missing.append(name.strip());continue
            found.append(name.strip())
            for month in range(1,13):
                text=str(sheet.cell(row,month+1).value or '')
                raw=DATE_RE.findall(text)
                if not raw:continue
                if len(raw)!=2:raise VacationError(f'В строке «{name}», {MONTHS[month-1]} должны быть две даты')
                start,end=('%04d-%02d-%02d'%(int(y),int(m),int(d)) for d,m,y in raw)
                dates(year,month,start,end)
                lower=sheet.cell(row+14,month+1).value
                help_flag=bool(lower and re.search(r'\b[МM]\b',str(lower).upper()))
                entries.append((uid,year,month,start,end,int(help_flag)))
        if not mapping:raise VacationError('В файле не найдены строки сотрудников')
        keys=[(entry[0],entry[2]) for entry in entries]
        if len(set(keys))!=len(keys):raise VacationError('Две строки файла сопоставлены одному сотруднику и месяцу; исправьте соответствия')
        return {'entries':entries,'matched':found,'unmatched':missing,'warnings':warnings,'mapping':mapping}
    finally:book.close()

def export_xlsx(year,people,rows):
    from openpyxl import Workbook
    from openpyxl.styles import Font,PatternFill,Alignment,Border,Side
    from openpyxl.utils import get_column_letter
    book=Workbook();sheet=book.active;sheet.title=f'Отпуск {year}'
    sheet.sheet_view.showGridLines=False
    sheet.column_dimensions['A'].width=29.57  # Source sheet: compact FIO column with wrapping.
    sheet.column_dimensions['B'].width=11.14
    for i in range(3,14):sheet.column_dimensions[get_column_letter(i)].width=13
    sheet.column_dimensions['N'].width=12.43
    line=Side(style='thin',color='CDE2DD')
    border=Border(left=line,right=line,top=line,bottom=line)
    cells={(r['user_id'],r['month']):r for r in rows}
    for offset in (0,len(people)+3):
        row=offset+1
        sheet.merge_cells(start_row=row,start_column=1,end_row=row+1,end_column=1)
        sheet.cell(row,1,'ФИО')
        sheet.merge_cells(start_row=row,start_column=2,end_row=row,end_column=13)
        sheet.cell(row,2,year)
        for month in range(1,13):sheet.cell(row+1,month+1,MONTHS[month-1])
        if offset==0:sheet.cell(row+1,14,'Итого дней')
        for cell in sheet[row]+sheet[row+1]:
            cell.fill=PatternFill('solid',fgColor='DDEFEA');cell.font=Font(bold=True,color='205E54')
            cell.alignment=Alignment(horizontal='center',vertical='center',wrap_text=True)
        for index,person in enumerate(people):
            at=row+2+index;sheet.cell(at,1,person['name']).alignment=Alignment(horizontal='center',vertical='center',wrap_text=True)
            days=0
            for month in range(1,13):
                cell=sheet.cell(at,month+1);entry=cells.get((person['id'],month))
                if entry:
                    if offset==0:
                        cell.value=dt.date.fromisoformat(entry['start_date']).strftime('%d.%m.%Y')+'\n'+dt.date.fromisoformat(entry['end_date']).strftime('%d.%m.%Y')
                        cell.fill=PatternFill('solid',fgColor='E9F5F1');days+=entry['days']
                    else:
                        cell.value='+М' if entry['material_help'] else '+'
                        if entry['material_help']:cell.fill=PatternFill('solid',fgColor='D7F3DC')
                cell.alignment=Alignment(horizontal='center',vertical='center',wrap_text=True)
            if offset==0:sheet.cell(at,14,days or None)
            sheet.row_dimensions[at].height=32 if offset==0 else 30
        for month in range(1,13):
            n=calendar.monthrange(year,month)[1]
            sheet.cell(row+2+len(people),month+1,f'{n} дней в месяце')
        sheet.row_dimensions[row].height=27;sheet.row_dimensions[row+1].height=28
    for row in sheet:
        for cell in row:
            if cell.__class__.__name__!='MergedCell':cell.border=border
    sheet.freeze_panes='B3'
    sheet.print_options.horizontalCentered=True
    sheet.sheet_properties.pageSetUpPr.fitToPage=True
    sheet.page_setup.orientation='landscape';sheet.page_setup.paperSize=sheet.PAPERSIZE_A3
    sheet.page_setup.fitToWidth=1;sheet.page_setup.fitToHeight=1
    sheet.print_area=f'A1:N{2*len(people)+6}'
    output=io.BytesIO();book.save(output);return output.getvalue()
