#!/usr/bin/env python3
"""Первичная загрузка исходного Excel. Не трогает уже созданные данные."""
import datetime as dt
import secrets
import warnings
from pathlib import Path
import openpyxl
from app import connect,init_db,hashed_password,TODAY

ROOT=Path(__file__).resolve().parent
USERNAMES=['gorodinskiy','goroh','gradusova','ksenevich','koshel','novik','savich','kotik','furs','omelchuk']

def category(title,number):
    if number>=18: return 'ПК УРГ'
    if 'Панорам' in title or 'карт' in title or 'фитинг' in title or 'подвал' in title: return 'МПК Панорама'
    return '7ЭГ и документация'

def run():
    init_db()
    with connect() as db:
        if db.execute('SELECT COUNT(*) FROM users').fetchone()[0]:
            print('База уже заполнена. Импорт пропущен.');return
        with warnings.catch_warnings():
            warnings.simplefilter('ignore',UserWarning)
            book=openpyxl.load_workbook(ROOT/'source.xlsx',data_only=True)
        report=book['Отчёт']
        credentials=['ПЕРВИЧНЫЕ ДАННЫЕ ДЛЯ ВХОДА В FORMA','', 'Смените пароли после первого входа. Не публикуйте этот файл.','']
        password='admin'
        db.execute('INSERT INTO users(name,username,password_hash,role,is_staff) VALUES(?,?,?,?,0)',('Администратор','admin',hashed_password(password),'admin'))
        credentials.extend([f'Администратор: admin / {password}','', 'Сотрудники:'])
        user_ids=[]
        for sheet,username in zip(book.worksheets[:-1],USERNAMES):
            password=secrets.token_urlsafe(12)
            row=db.execute('INSERT INTO users(name,username,password_hash,role) VALUES(?,?,?,?)',(sheet.title,username,hashed_password(password),'employee'))
            user_ids.append(row.lastrowid)
            credentials.append(f'{sheet.title}: {username} / {password}')
        tasks=[]
        for number in range(23):
            rr=4+number; title=' '.join(str(report.cell(rr,1).value or '').split())
            unit=str(report.cell(rr,2).value or 'шт.').strip(); norm=float(report.cell(rr,14).value or 0)
            res=db.execute('INSERT INTO tasks(title,unit,norm,category) VALUES(?,?,?,?)',(title,unit,norm,category(title,number)))
            tasks.append(res.lastrowid)
        count=0
        for sheet,user_id in zip(book.worksheets[:-1],user_ids):
            dates={col:sheet.cell(1,col).value for col in range(3,sheet.max_column+1)}
            for number,task_id in enumerate(tasks):
                for col,date in dates.items():
                    if not isinstance(date,dt.datetime) or date.date()>TODAY():continue
                    val=sheet.cell(number+2,col).value
                    if isinstance(val,(int,float)) and not isinstance(val,bool) and val>0:
                        db.execute('INSERT OR IGNORE INTO entries(user_id,task_id,work_date,quantity,note) VALUES(?,?,?,?,?)',
                                   (user_id,task_id,date.date().isoformat(),val,''))
                        count+=1
        db.execute('INSERT INTO activity(actor_id,action,detail) VALUES(NULL,?,?)',('import',f'Импортировано {len(user_ids)} сотрудников, {len(tasks)} видов работ и {count} ежедневных записей из Excel'))
        db.commit()
    file=ROOT/'initial_credentials.txt';file.write_text('\n'.join(credentials)+'\n',encoding='utf-8');file.chmod(0o600)
    print(f'Импортировано: {len(user_ids)} пользователей, {len(tasks)} работ, {count} отметок. Пароли: {file}')

if __name__=='__main__':run()
