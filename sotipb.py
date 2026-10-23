"""Строгое преобразование «лист 1» → оформленный «свод 1» без выдумывания значений.

Загруженные пары файлов хранятся в D1/SQLite через таблицы истории приложения.
"""
import datetime as dt
import io
import re
import zipfile
from openpyxl import Workbook, load_workbook
from copy import copy
from pathlib import Path
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.comments import Comment
from openpyxl.utils import get_column_letter

HEADERS=[
 'Дата внесения изменений','Архивный номер','Новый архивный','Адрес',
 'Диаметр было/стало','Протяженность было/стало','Давление было/стало',
 'Вид прокладки было/стало','Балансовая принадлежность было/стало',
 'Заводской номер, номер проекта','Год строительства',
 'Наименование организации изготовителя ПОО','Кто вносил',
 'Причина корректировки','Передано в ТИ',
 'Тип ПОО для свода (уточнение)','L, м для свода (уточнение)',
 'D, мм для свода (уточнение)']
BASE=15
SUMMARY_LAYOUT=Path(__file__).with_name('sotipb-summary-layout.xlsx')
MAX_INPUT=3*1024*1024
MAX_ROWS=5000
MAX_EXPANDED=30*1024*1024

class ConversionError(Exception): pass

def _s(value):
    if value is None:return ''
    if isinstance(value,(dt.datetime,dt.date)):return value.strftime('%d.%m.%Y')
    if isinstance(value,float) and value.is_integer():return str(int(value))
    return str(value).strip()

def _cell(sheet,row,col):
    cell=sheet.cell(row,col)
    if cell.data_type=='f':
        raise ConversionError(f'Формула в строке {row}, столбец {get_column_letter(col)}: замените её значением')
    return cell.value

def _split(text):
    # Slash denotes old/new only in change columns; project IDs contain slashes.
    return [p.strip() for p in text.split('/',1)] if '/' in text else [text.strip()]

def _material(text):
    t=text.casefold().replace('ё','е')
    if re.search(r'пэ|полиэтилен|\bпе\b',t):return 'полиэтиленовый'
    if re.search(r'сталь|стальн|ст\b|ст$',t):return 'стальной'
    return None

def _pressure(text):
    t=re.sub(r'[\s.\-]+','',text.casefold().replace('ё','е'))
    if t in ('нд','низкое','низкогодавления'):return 'низкого давления'
    if t in ('сд','среднее','среднегодавления'):return 'среднего давления'
    if t in ('вд1','в1','высокое1','высокого давления i категории'.replace(' ','')):
        return 'высокого давления I категории'
    if t in ('вд2','в2','высокое2','высокого давления ii категории'.replace(' ','')):
        return 'высокого давления II категории'
    return None

def _type(row,group,warnings):
    if row[15]:return row[15]
    parts=_split(row[4])
    pressure=_split(row[6])
    types=[];inferred=False
    for i,diam in enumerate(parts):
        mat=_material(diam) or next((_material(r[4]) for r in group if _material(r[4])),None)
        if not mat and diam:
            mat='стальной'
            if not inferred:warnings.append(f'Строка {row.number}: у диаметра не указан материал; по образцу принято «стальной» (уточнение в P)')
            inferred=True
        p=_pressure(pressure[min(i,len(pressure)-1)])
        if mat and p:types.append('Газопровод '+mat+' '+p)
        else:
            warnings.append(f'Строка {row.number}: тип ПОО не определён; укажите материал и давление или уточнение в P')
            return ''
    return types[0] if len(types)==1 or len(set(types))==1 else '\n'.join(f'{t} ({"было" if i==0 else "стало"})' for i,t in enumerate(types))

def _length(row,warnings):
    # Existing Q is a fallback only for an empty F; explicit Q content is
    # already a ready-to-use summary value and is not reformatted.
    if not row[5]:return row.raw[16] if row[16] else ''
    text=row[5]
    if '/' not in text:return row.raw[5]
    if text.count('/')!=1 or not all(part.strip() for part in text.split('/')):
        warnings.append(f'Строка {row.number}: F содержит неоднозначное разделение «/»; значение оставлено без преобразования')
        return row.raw[5]
    old,new=(part.strip() for part in text.split('/'))
    return f'{old} (было)\n{new} (стало)'

_DIAMETER_NUMBER=re.compile(
    r'\s*(?:[dDØø⌀Фф]\s*|Ду\s*)?(\d+(?:[.,]\d+)?)\s*'
    r'(?:мм\.?|mm\.?|пэ|пе|ст|сталь|стальной|полиэтилен(?:овый)?)?\s*',re.I)

def _diameter(row,warnings):
    # Existing R is a ready-to-use fallback only for an empty E.
    if not row[4]:return row.raw[17] if row[17] else ''
    text=row[4]
    if text.count('/')>1:
        warnings.append(f'Строка {row.number}: в E более двух диаметров; D свода оставлен пустым, уточните E')
        return ''
    parts=[part.strip() for part in text.split('/')]
    if not all(parts):
        warnings.append(f'Строка {row.number}: неполное значение E; D свода оставлен пустым, уточните E')
        return ''
    numbers=[]
    for part in parts:
        match=_DIAMETER_NUMBER.fullmatch(part)
        if not match:
            warnings.append(f'Строка {row.number}: диаметр «{part[:40]}» не является однозначным числом; D свода оставлен пустым, уточните E')
            return ''
        numbers.append(match.group(1))
    layers=[part.strip() for part in row[7].split('/')]
    if len(layers)>2 or len(layers)==2 and len(numbers)!=2:
        warnings.append(f'Строка {row.number}: значения H не соответствуют E; проверьте вид прокладки')
    def label(number,index):
        layer=layers[min(index,len(layers)-1)] if layers else ''
        return (number+' '+layer).strip()
    if len(numbers)==1:return label(numbers[0],0)
    return label(numbers[0],0)+' (было)\n'+label(numbers[1],1)+' (стало)'

def _year(text):return text[:-2] if text.endswith('.0') and text[:-2].isdigit() else text

def _safe_write(cell,value):
    # Never turn uploaded text into an Excel formula (including CSV-style prefixes).
    cell.value=value
    if isinstance(value,str) and value.startswith(('=','+','-','@','\t','\r')):
        cell.data_type='s'

def _source_style(sheet):
    sheet.freeze_panes='D2';sheet.sheet_view.showGridLines=False
    widths=[20,20,20,63,22,24,22,23,27,34,20,39,19,47,19,35,30,30]
    for i,w in enumerate(widths,1):sheet.column_dimensions[get_column_letter(i)].width=w
    for c,title in enumerate(HEADERS,1):
        cell=sheet.cell(1,c);cell.value=title
        cell.font=Font(name='Arial',size=11,bold=True,color='FFFFFF')
        cell.fill=PatternFill('solid',fgColor='1B544D')
        cell.alignment=Alignment(wrap_text=True,vertical='center')
        if c>BASE:
            cell.comment=Comment('Необязательное уточнение. Q/R используются только если соответственно F/E в этой строке пусты; P уточняет тип ПОО.','Forma')
    sheet.row_dimensions[1].height=56
    sheet.auto_filter.ref='A1:R1'
    sheet.print_options.horizontalCentered=True
    sheet.sheet_properties.pageSetUpPr.fitToPage=True
    sheet.page_setup.fitToWidth=1

def template():
    book=Workbook();sheet=book.active;sheet.title='лист 1';_source_style(sheet)
    hints={2:'Архивный номер, если C пуст. Например: 5-65/2',3:'Новый архивный номер: при наличии он имеет приоритет над B',4:'Адрес ПОО',5:'Число + вид прокладки H в D свода; для 600/630 — две строки с было/стало',6:'Для 3.3/0.2 в C свода — две строки: 3.3 (было), 0.2 (стало)',7:'нд, сд, вд1 или В2',8:'подземный или надземный'}
    for col,example in hints.items():sheet.cell(1,col).comment=Comment('Вводите данные начиная со строки 2. '+example+'; пустые ячейки допускаются, но в своде не будут заполняться догадками.','Forma')
    stream=io.BytesIO();book.save(stream);book.close();return stream.getvalue()

def convert(blob):
    if not blob.startswith(b'PK\x03\x04') or len(blob)>MAX_INPUT:raise ConversionError('Нужен XLSX размером до 3 МБ')
    try:
        with zipfile.ZipFile(io.BytesIO(blob)) as z:
            if len(z.namelist())>300 or sum(x.file_size for x in z.infolist())>MAX_EXPANDED:raise ConversionError('Слишком большой XLSX после распаковки')
        book=load_workbook(io.BytesIO(blob),read_only=True,data_only=False,keep_links=False)
    except (zipfile.BadZipFile,ValueError,OSError) as exc:raise ConversionError('Файл не читается как XLSX') from exc
    try:
        source=book['лист 1'] if 'лист 1' in book else book.worksheets[0]
        first=[_s(_cell(source,1,i)) for i in range(1,BASE+1)]
        if first!=HEADERS[:BASE]:
            raise ConversionError('Заголовки A–O не совпадают с «лист 1». Скачайте болванку и перенесите свои строки в неё')
        extra_headers=[_s(_cell(source,1,i)) for i in range(BASE+1,len(HEADERS)+1)]
        if source.max_row and source.max_row>30000:raise ConversionError('В книге слишком много строк; уберите лишнее форматирование ниже таблицы')
        rows=[];warnings=[]
        for number,cells in enumerate(source.iter_rows(min_row=2,max_col=len(HEADERS)),2):
            values=[];raw=[]
            for i,cell in enumerate(cells):
                if cell.data_type=='f':raise ConversionError(f'Формула в строке {number}, столбец {get_column_letter(i+1)}: замените значением')
                value=cell.value
                values.append(_s(value));raw.append(value)
            if not any(values):continue
            if not any(values[:BASE]):continue
            if any(values[BASE:]) and extra_headers!=HEADERS[BASE:]:
                raise ConversionError(f'Строка {number}: значения в P–R допустимы только под заголовками болванки СОТИПБ')
            if not values[3]:warnings.append(f'Строка {number}: адрес пуст')
            if not values[4] and not values[17]:warnings.append(f'Строка {number}: диаметр пуст')
            if not values[5] and not values[16]:warnings.append(f'Строка {number}: длина отсутствует и не будет придумана')
            if not values[6]:warnings.append(f'Строка {number}: давление пусто')
            if not values[7]:warnings.append(f'Строка {number}: вид прокладки пуст')
            row=Record(values,number,raw)
            rows.append(row)
            if len(rows)>MAX_ROWS:raise ConversionError('Свыше 5000 строк данных: разделите исходный файл на части')
        if not rows:raise ConversionError('В «лист 1» нет заполненных строк под заголовком')
        if not SUMMARY_LAYOUT.is_file():
            raise ConversionError('Не опубликован файл оформления sotipb-summary-layout.xlsx. Загрузите его в корень GitHub')
        result=load_workbook(SUMMARY_LAYOUT);input_sheet=result.create_sheet('лист 1',0);_source_style(input_sheet)
        for index,row in enumerate(rows,2):
            for col,text in enumerate(row.raw,1):
                cell=input_sheet.cell(index,col);_safe_write(cell,text)
                if isinstance(text,(dt.datetime,dt.date)):cell.number_format='DD.MM.YYYY'
            input_sheet.row_dimensions[index].height=30
        input_sheet.auto_filter.ref=f'A1:R{len(rows)+1}'
        summary=result['свод 1']
        body_styles=[copy(summary.cell(5,c)._style) for c in range(1,11)]
        groups=[]
        for row in rows:
            # C (new archive number) wins; B is used only when C is empty.
            # Group by the same effective number that appears in the summary.
            key=((row[2] or row[1]).casefold(),row[3].casefold())
            if groups and groups[-1][0]==key:groups[-1][1].append(row)
            else:groups.append((key,[row]))
        line=5
        for _,group in groups:
            start=line
            for row in group:
                dtype=_type(row,group,warnings)
                address=row[3]
                number=row[2] or row[1]
                if number:address+='\n('+('нов. арх. № ' if row[2] else 'арх. № ')+number+')'
                # F splits into labeled old/new lengths. E supplies numeric
                # diameters, with the corresponding laying method from H.
                # Q/R remain fallback values only when F/E are empty.
                length=_length(row,warnings)
                diameter=_diameter(row,warnings)
                year=row.raw[10] if isinstance(row.raw[10],(int,float)) and not isinstance(row.raw[10],bool) else _year(row[10])
                output=[dtype,address,length,diameter,row[9],year,row[11],'','','']
                for col,text in enumerate(output,1):
                    cell=summary.cell(line,col);_safe_write(cell,text)
                    # The actual sample's unfilled, black medium-border cells:
                    # copy the clean visual template, not the application's green UI style.
                    cell._style=copy(body_styles[col-1])
                    if isinstance(text,str) and '\n' in text and not cell.alignment.wrap_text:
                        aligned=copy(cell.alignment);aligned.wrap_text=True;cell.alignment=aligned
                # Ensure both lines are visible, not merely present in the XLSX XML.
                visible_lines=max((value.count('\n')+1 for value in output if isinstance(value,str)),default=1)
                if visible_lines>1:
                    current=summary.row_dimensions[line].height or summary.sheet_format.defaultRowHeight or 15
                    summary.row_dimensions[line].height=max(current,visible_lines*16+5)
                line+=1
            end=line-1
            # Merge only shared, identical group values; preserve differing types and
            # diameter/length per input row. H–J remain empty as in supplied example.
            for col in (2,5,6,7,8,9,10):
                values=[summary.cell(r,col).value for r in range(start,end+1)]
                if end>start and len(set(values))==1:summary.merge_cells(start_row=start,start_column=col,end_row=end,end_column=col)
            run=start
            while run<=end:
                stop=run
                while stop<end and summary.cell(stop+1,1).value==summary.cell(run,1).value:stop+=1
                if stop>run and summary.cell(run,1).value:summary.merge_cells(start_row=run,start_column=1,end_row=stop,end_column=1)
                run=stop+1
        # Leave print settings, widths, default row heights and header merges
        # exactly as in the supplied summary sheet.
        stream=io.BytesIO();result.save(stream);result.close()
        return stream.getvalue(), {'rows':len(rows),'groups':len(groups),'warnings':warnings[:100], 'warning_count':len(warnings)}
    finally:book.close()

class Record(list):
    def __init__(self,values,number,raw):super().__init__(values);self.number=number;self.raw=raw
